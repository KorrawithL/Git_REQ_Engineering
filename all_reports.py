import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import io
import time
import json
import math
import calendar 
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import get_db_connection, log_activity
from config import branch_dict

# =========================================================================
# 🌟 ท่าไม้ตายล้างกราฟค้าง (Force UI Flush) 🌟
# =========================================================================
ui_flusher = st.empty()
with ui_flusher:
    st.markdown("<div style='height: 1px;'></div>", unsafe_allow_html=True)
time.sleep(0.05) 
ui_flusher.empty() 
# =========================================================================

# ==========================================================
# 🛑 1. ฟังก์ชัน Popup ยืนยันการลบข้อมูล (Delete Dialogs)
# ==========================================================
@st.dialog("⚠️ ยืนยันการลบข้อมูล (Tab 1)")
def delete_record_dialog_t1(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.error(f"คุณแน่ใจหรือไม่ว่าต้องการลบข้อมูล ID: {rec_id} ?")
    st.write(f"**วันที่:** {row_data.get('record_date')} | **สาขา:** {row_data.get('branch_name')}")
    st.write(f"**เครื่องจักร:** {row_data.get('machine_name')}")
    st.markdown("<span style='color:#EF4444; font-size:14px;'>* การกระทำนี้ไม่สามารถกู้คืนข้อมูลได้</span>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        if st.button("❌ ยกเลิก", use_container_width=True, key=f"cancel_d1_{rec_id}"): st.rerun()
    with c2:
        if st.button("🗑️ ยืนยันการลบ", type="primary", use_container_width=True, key=f"confirm_d1_{rec_id}"):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"DELETE FROM machine_trans WHERE {pk_col}=%s", (rec_id,))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "DELETE", "All Report: Tab 1", f"ลบข้อมูล ID: {rec_id}")
                st.success("🗑️ ลบข้อมูลเรียบร้อยแล้ว!")
                time.sleep(1)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("⚠️ ยืนยันการลบข้อมูล (Tab 2)")
def delete_record_dialog_t2(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.error(f"คุณแน่ใจหรือไม่ว่าต้องการลบข้อมูล ID: {rec_id} ?")
    st.write(f"**วันที่:** {row_data.get('record_date')} | **สาขา:** {row_data.get('branch_name')}")
    st.write(f"**ทะเบียนรถ:** {row_data.get('engine_code')} ({row_data.get('type_name')})")
    st.markdown("<span style='color:#EF4444; font-size:14px;'>* การกระทำนี้ไม่สามารถกู้คืนข้อมูลได้</span>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        if st.button("❌ ยกเลิก", use_container_width=True, key=f"cancel_d2_{rec_id}"): st.rerun()
    with c2:
        if st.button("🗑️ ยืนยันการลบ", type="primary", use_container_width=True, key=f"confirm_d2_{rec_id}"):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"DELETE FROM fuel_records WHERE {pk_col}=%s", (rec_id,))
                    conn.commit()

                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "DELETE", "All Report: Tab 2", f"ลบข้อมูล ID: {rec_id}")
                st.success("🗑️ ลบข้อมูลเรียบร้อยแล้ว!")
                time.sleep(1)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")
            finally:
                if conn:
                    conn.close()

@st.dialog("⚠️ ยืนยันการลบข้อมูล (Tab 3)")
def delete_record_dialog_t3(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.error(f"คุณแน่ใจหรือไม่ว่าต้องการลบข้อมูล ID: {rec_id} ?")
    st.write(f"**วันที่:** {row_data.get('record_date')} | **สาขา:** {row_data.get('branch_name')}")
    st.write(f"**จำนวนตรวจเช็คทั้งหมด:** {row_data.get('total_count')} ครั้ง | **ตก:** {row_data.get('total_drop')} ครั้ง")
    st.markdown("<span style='color:#EF4444; font-size:14px;'>* การกระทำนี้ไม่สามารถกู้คืนข้อมูลได้</span>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        if st.button("❌ ยกเลิก", use_container_width=True, key=f"cancel_d3_{rec_id}"): st.rerun()
    with c2:
        if st.button("🗑️ ยืนยันการลบ", type="primary", use_container_width=True, key=f"confirm_d3_{rec_id}"):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"DELETE FROM boiler_pressure_records WHERE {pk_col}=%s", (rec_id,))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "DELETE", "All Report: Tab 3", f"ลบข้อมูล ID: {rec_id}")
                st.success("🗑️ ลบข้อมูลเรียบร้อยแล้ว!")
                time.sleep(1)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("⚠️ ยืนยันการลบข้อมูล (Tab 4)")
def delete_record_dialog_t4(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.error(f"คุณแน่ใจหรือไม่ว่าต้องการลบข้อมูล ID: {rec_id} ?")
    st.write(f"**วันที่:** {row_data.get('record_date')} | **สาขา:** {row_data.get('branch_name')}")
    st.write(f"**ผลิตไอน้ำ:** {row_data.get('steam_production')} ตัน | **ชม.ทำงาน:** {row_data.get('working_hours')} ชม.")
    st.markdown("<span style='color:#EF4444; font-size:14px;'>* การกระทำนี้ไม่สามารถกู้คืนข้อมูลได้</span>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        if st.button("❌ ยกเลิก", use_container_width=True, key=f"cancel_d4_{rec_id}"): st.rerun()
    with c2:
        if st.button("🗑️ ยืนยันการลบ", type="primary", use_container_width=True, key=f"confirm_d4_{rec_id}"):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"DELETE FROM boiler_fuel_records WHERE {pk_col}=%s", (rec_id,))
                    conn.commit()

                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "DELETE", "All Report: Tab 4", f"ลบข้อมูล ID: {rec_id}")
                st.success("🗑️ ลบข้อมูลเรียบร้อยแล้ว!")
                time.sleep(1)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")
            finally:
                if conn:
                    conn.close()

@st.dialog("⚠️ ยืนยันการลบข้อมูล (ยอดรวมสาขา)")
def delete_record_dialog_oven(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.error(f"คุณแน่ใจหรือไม่ว่าต้องการลบข้อมูล ID: {rec_id} ?")
    st.write(f"**วันที่:** {row_data.get('record_date')} | **สาขา:** {row_data.get('branch_name')}")
    st.write(f"**จำนวนเตา:** {row_data.get('oven_qty')} เตา | **ไม้อบออก:** {row_data.get('wood_out_cubft')} ลบ.ฟ.")
    st.markdown("<span style='color:#EF4444; font-size:14px;'>* การกระทำนี้ไม่สามารถกู้คืนข้อมูลได้</span>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        if st.button("❌ ยกเลิก", use_container_width=True, key=f"c_d_ov_{rec_id}"): st.rerun()
    with c2:
        if st.button("🗑️ ยืนยันการลบ", type="primary", use_container_width=True, key=f"d_d_ov_{rec_id}"):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"DELETE FROM daily_wood_oven_records WHERE {pk_col}=%s", (rec_id,))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "DELETE", "All Report: Tab 4 Oven", f"ลบข้อมูล ID: {rec_id}")
                st.success("🗑️ ลบข้อมูลเรียบร้อยแล้ว!")
                time.sleep(1)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

# ==========================================================
# 🛠️ 2. ฟังก์ชัน Popup สำหรับแก้ไขข้อมูล (Update Dialogs)
# ==========================================================
@st.dialog("🛠️ แก้ไขข้อมูล (Tab 1)", width="large")
def update_record_dialog_t1(row_data, pk_col):
    rec_id = row_data[pk_col]
    curr_m = str(row_data.get('machine_name') or 'ไม่ทราบชื่อเครื่องจักร')
    curr_branch = str(row_data.get('branch_name') or '')
    curr_branch_id = row_data.get('branch_id')
    curr_m_id = row_data.get('machine_id')
    
    st.info(f"ID: {rec_id} | เครื่องเดิม: {curr_m} ({curr_branch})")
    
    raw_role = st.session_state.get('role_tab') or st.session_state.get('role') or 'user'
    current_role = str(raw_role).strip().lower()
    
    new_branch_id = curr_branch_id
    new_m_id = curr_m_id
    btn_disabled = False
    
    c_top1, c_top2 = st.columns([1.5, 1.5])
    
    if current_role == "admin":
        edit_branch_options = list(branch_dict.keys())
        def_branch_lbl = next((k for k, v in branch_dict.items() if v == curr_branch_id), edit_branch_options[0] if edit_branch_options else None)
        def_branch_idx = edit_branch_options.index(def_branch_lbl) if def_branch_lbl in edit_branch_options else 0
        
        with c_top1:
            e_branch_lbl = st.selectbox("🏢 1. แก้ไขสาขา *", edit_branch_options, index=def_branch_idx, key=f"e1_br_{rec_id}")
            new_branch_id = branch_dict[e_branch_lbl]
        
        m_dict = {}
        try:
            conn = get_db_connection()
            with conn.cursor() as cur:
                cur.execute("SELECT id, CONVERT(machine_name USING utf8mb4) AS m_name FROM machines_set WHERE branch_id = %s AND machine_type = 'machine' AND is_active = 1", (new_branch_id,))
                for m in cur.fetchall(): m_dict[m['m_name']] = m['id']
            conn.close()
        except Exception: pass
        
        m_options = list(m_dict.keys())
        if not m_options: m_options = ["-- ไม่พบเครื่องจักร --"]
        
        def_m_idx = 0
        if new_branch_id == curr_branch_id and curr_m in m_options: def_m_idx = m_options.index(curr_m)
        with c_top2:
            e_machine_lbl = st.selectbox("⚙️ 2. แก้ไขเครื่องจักร *", m_options, index=def_m_idx, key=f"e1_m_{rec_id}")
            if e_machine_lbl != "-- ไม่พบเครื่องจักร --": new_m_id = m_dict.get(e_machine_lbl, curr_m_id)
            else: btn_disabled = True
    else:
        with c_top1: st.text_input("🏢 1. สาขา (ไม่อนุญาตให้แก้ไข)", value=curr_branch, disabled=True, key=f"e1_br_{rec_id}")
        with c_top2: st.text_input("⚙️ 2. เครื่องจักร (ไม่อนุญาตให้แก้ไข)", value=curr_m, disabled=True, key=f"e1_m_{rec_id}")
    
    st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        e_date = st.date_input("แก้ไข วันที่", format="DD/MM/YYYY", value=pd.to_datetime(row_data.get('record_date')), max_value=pd.to_datetime("today"))
        e_qty = st.number_input("แก้ไข จำนวนเครื่องจักร", value=int(row_data.get('machine_qty') or 1), min_value=1)
    with col2:
        e_work = st.number_input("แก้ไข ชม.ทำงาน", value=float(row_data.get('working_hours') or 0.0), min_value=0.0, format="%.2f")
        e_break = st.number_input("แก้ไข ชม.เบรกดาวน์", value=float(row_data.get('breakdown_hours') or 0.0), min_value=0.0, format="%.2f")
    e_rem = st.text_area("แก้ไข หมายเหตุ", value=str(row_data.get('remarks') or ''))
    
    if st.button("💾 บันทึกการแก้ไข", use_container_width=True, type="primary", disabled=btn_disabled):
        if (e_work + e_break) > 24.0:
            st.error("⚠️ ไม่สามารถบันทึกได้: ชั่วโมงทำงานรวมกับเบรกดาวน์ ต้องไม่เกิน 24 ชั่วโมงต่อวันครับ")
            return
        with st.spinner("กำลังบันทึกข้อมูล..."):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"UPDATE machine_trans SET branch_id=%s, machine_id=%s, record_date=%s, machine_qty=%s, working_hours=%s, breakdown_hours=%s, remarks=%s, updated_at=NOW() WHERE {pk_col}=%s",
                                (new_branch_id, new_m_id, e_date, e_qty, e_work, e_break, e_rem, rec_id))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "UPDATE", "All Report: Tab 1", f"แก้ไข ID: {rec_id}")
                st.toast("✅ อัปเดตข้อมูลสำเร็จ!", icon="🎉") 
                time.sleep(1.2)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("🛠️ แก้ไขข้อมูล (Tab 2)", width="large")
def update_record_dialog_t2(row_data, pk_col, engine_lbl_list, engine_details_t2):
    rec_id = row_data[pk_col]
    curr_branch = str(row_data.get('branch_name') or '')
    curr_type = str(row_data.get('type_name') or '')
    curr_code = str(row_data.get('engine_code') or '')
    curr_branch_id = row_data.get('branch_id')

    st.info(f"ID: {rec_id} | ทะเบียนเดิม: {curr_code} ({curr_branch})")
    
    raw_role = st.session_state.get('role_tab') or st.session_state.get('role') or 'user'
    current_role = str(raw_role).strip().lower()
    c_sel1, c_sel2, c_sel3 = st.columns(3)
    
    new_type, new_code, inputs_disabled, e_branch = curr_type, curr_code, False, curr_branch 

    if current_role == "admin":
        all_engines = []
        try:
            conn = get_db_connection()
            with conn.cursor() as cur:
                sql = """SELECT CONVERT(e.machine_name USING utf8mb4) AS engine_code, CONVERT(et.type_name USING utf8mb4) AS type_name, CONVERT(b.branch_name USING utf8mb4) AS branch_name FROM machines_set e LEFT JOIN engine_types et ON e.engine_type_id = et.id LEFT JOIN branches b ON e.branch_id = b.id WHERE e.machine_type = 'engine' AND e.is_active = 1"""
                cur.execute(sql)
                all_engines = cur.fetchall()
            conn.close()
        except Exception: pass

        available_branches = sorted(list(set([e['branch_name'] for e in all_engines if e['branch_name']])))
        def_branch_idx = available_branches.index(curr_branch) + 1 if curr_branch in available_branches else 0
        e_branch = c_sel1.selectbox("🏢 1. แก้ไขสาขา *", ["-- เลือกสาขา --"] + available_branches, index=def_branch_idx, key=f"e2_br_{rec_id}")

        f2_type_disabled = (e_branch == "-- เลือกสาขา --")
        if not f2_type_disabled:
            filtered_by_branch = [e for e in all_engines if e['branch_name'] == e_branch]
            available_types = sorted(list(set([e['type_name'] for e in filtered_by_branch if e['type_name']])))
        else: available_types, filtered_by_branch = [], []
            
        def_type_idx = available_types.index(curr_type) + 1 if (curr_type in available_types and e_branch == curr_branch) else 0
        e_type = c_sel2.selectbox("🚜 2. แก้ไขชนิดรถ *", ["-- เลือกประเภท --"] + available_types, index=def_type_idx, key=f"e2_ty_{rec_id}", disabled=f2_type_disabled)

        f2_code_disabled = (e_type == "-- เลือกประเภท --") or f2_type_disabled
        if not f2_code_disabled:
            filtered_by_type = [e for e in filtered_by_branch if e['type_name'] == e_type]
            available_codes = sorted(list(set([e['engine_code'] for e in filtered_by_type if e['engine_code']])))
        else: available_codes, filtered_by_type = [], []
            
        def_code_idx = available_codes.index(curr_code) + 1 if (curr_code in available_codes and e_type == curr_type and e_branch == curr_branch) else 0
        e_code = c_sel3.selectbox("🏷️ 3. แก้ไขทะเบียนรถ *", ["-- เลือกรถ --"] + available_codes, index=def_code_idx, key=f"e2_cd_{rec_id}", disabled=f2_code_disabled)
        
        inputs_disabled = (e_code == "-- เลือกรถ --") or f2_code_disabled
        if not inputs_disabled: new_type, new_code = e_type, e_code
    else:
        with c_sel1: st.text_input("🏢 1. สาขา (ไม่อนุญาตให้แก้ไข)", value=curr_branch, disabled=True, key=f"e2_br_{rec_id}")
        with c_sel2: st.text_input("🚜 2. ชนิด / ประเภทรถ", value=curr_type, disabled=True, key=f"e2_ty_{rec_id}")
        with c_sel3: st.text_input("🏷️ 3. ทะเบียนรถ", value=curr_code, disabled=True, key=f"e2_cd_{rec_id}")

    st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        e_date = st.date_input("วันที่", value=pd.to_datetime(row_data.get('record_date')), format="DD/MM/YYYY", max_value=pd.to_datetime("today"), key=f"e2_dt_{rec_id}")
        e_liters = st.number_input("ปริมาณน้ำมัน (ลิตร) *", min_value=0.0, value=float(row_data.get('fuel_liters') or 0.0), key=f"e2_lt_{rec_id}")
    with col2:
        e_work = st.number_input("จำนวนชั่วโมงทำงาน (ชม.) *", min_value=0.0, step=0.5, value=float(row_data.get('working_hours') or 0.0), key=f"e2_wk_{rec_id}")
        e_rem = st.text_area("หมายเหตุ", value=str(row_data.get('remark') or ''), key=f"e2_rm_{rec_id}")

    if st.button("💾 บันทึกการแก้ไข", use_container_width=True, type="primary", disabled=inputs_disabled, key=f"btn_save_e2_{rec_id}"):
        if e_work > 24.0:
            st.error("⚠️ ไม่สามารถบันทึกได้: ชั่วโมงทำงานรถยนต์ต้องไม่เกิน 24 ชั่วโมงต่อวัน")
            return
        with st.spinner("กำลังบันทึกข้อมูล..."):
            try:
                u_liter_hr = round(e_liters / e_work, 2) if e_work > 0 else 0.00
                new_branch_id = curr_branch_id
                conn = get_db_connection()
                with conn.cursor() as cur:
                    if current_role == "admin":
                        cur.execute("SELECT id FROM branches WHERE branch_name = %s", (e_branch,))
                        b_res = cur.fetchone()
                        if b_res: new_branch_id = b_res['id']
                    cur.execute(f"""UPDATE fuel_records SET branch_id=%s, record_date=%s, engine_code=CONVERT(%s USING utf8mb4), type_name=CONVERT(%s USING utf8mb4), fuel_liters=%s, working_hours=%s, liter_hr=%s, remark=CONVERT(%s USING utf8mb4) WHERE {pk_col}=%s""",
                                (new_branch_id, e_date, new_code, new_type, e_liters, e_work, u_liter_hr, e_rem, rec_id))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "UPDATE", "All Report: Tab 2", f"แก้ไข ID: {rec_id} ({new_code})")
                st.toast("✅ อัปเดตข้อมูลสำเร็จ!", icon="🚚")
                time.sleep(1.2)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("🛠️ แก้ไขข้อมูล (Tab 3)", width="large")
def update_record_dialog_t3(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.info(f"ID: {rec_id} | สาขาเดิม: {row_data.get('branch_name')}")
    
    raw_role = st.session_state.get('role_tab') or st.session_state.get('role') or 'user'
    user_allowed_branches = st.session_state.get('allowed_branches', [str(st.session_state.branch_id)])
    if str(raw_role).strip().lower() in ["admin", "reporter"]: edit_branch_options = list(branch_dict.keys())
    else: edit_branch_options = [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
    
    curr_branch = str(row_data.get('branch_name') or '')
    def_branch_idx = edit_branch_options.index(curr_branch) if curr_branch in edit_branch_options else 0
    e_branch_lbl = st.selectbox("🏢 1. แก้ไขสาขา *", edit_branch_options, index=def_branch_idx, key=f"e3_br_{rec_id}")
    new_branch_id = branch_dict[e_branch_lbl]
    
    st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        e_date = st.date_input("แก้ไข วันที่", format="DD/MM/YYYY", value=pd.to_datetime(row_data.get('record_date')), max_value=pd.to_datetime("today"))
        e_tot = st.number_input("แก้ไข จำนวนครั้งทั้งหมด", min_value=0, value=int(row_data.get('total_count') or 0))
    with col2:
        e_pm = st.number_input("แก้ไข ตกตาม PM", min_value=0, value=int(row_data.get('pm_drop') or 0))
        e_npm = st.number_input("แก้ไข ตกนอกเหนือ PM", min_value=0, value=int(row_data.get('non_pm_drop') or 0))
    e_rem = st.text_input("แก้ไข หมายเหตุ", value=str(row_data.get('remark') or ''))

    if st.button("💾 บันทึกการแก้ไข", use_container_width=True, type="primary"):
        if (e_pm + e_npm) > e_tot:
            st.error("⚠️ ไม่สามารถบันทึกได้: จำนวนครั้งที่ 'แรงดันตก' (PM + นอก PM) ต้องไม่มากกว่า 'จำนวนครั้งที่ตรวจเช็คทั้งหมด'")
            return
        with st.spinner("กำลังบันทึกข้อมูล..."):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"UPDATE boiler_pressure_records SET branch_id=%s, record_date=%s, total_count=%s, total_drop=%s, pm_drop=%s, non_pm_drop=%s, remark=%s WHERE {pk_col}=%s",
                                (new_branch_id, e_date, e_tot, (e_pm + e_npm), e_pm, e_npm, e_rem, rec_id))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "UPDATE", "All Report: Tab 3", f"แก้ไข ID: {rec_id}")
                st.toast("✅ อัปเดตข้อมูลสำเร็จ!", icon="💨")
                time.sleep(1.2)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("🛠️ แก้ไขข้อมูล (Tab 4)", width="large")
def update_record_dialog_t4(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.info(f"ID: {rec_id} | สาขาเดิม: {row_data.get('branch_name')}")
    
    raw_role = st.session_state.get('role_tab') or st.session_state.get('role') or 'user'
    user_allowed_branches = st.session_state.get('allowed_branches', [str(st.session_state.branch_id)])
    if str(raw_role).strip().lower() in ["admin", "reporter"]: edit_branch_options = list(branch_dict.keys())
    else: edit_branch_options = [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
    
    curr_branch = str(row_data.get('branch_name') or '')
    def_branch_idx = edit_branch_options.index(curr_branch) if curr_branch in edit_branch_options else 0
    
    c_top1, c_top2 = st.columns([1.5, 1.5])
    with c_top1:
        e_branch_lbl = st.selectbox("🏢 1. แก้ไขสาขา *", edit_branch_options, index=def_branch_idx, key=f"e4_br_{rec_id}")
        new_branch_id = branch_dict[e_branch_lbl]
    
    b_options = []
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            cur.execute("SELECT CONVERT(boiler_name USING utf8mb4) AS b_name FROM boilers WHERE branch_id = %s AND is_active = 1", (new_branch_id,))
            b_options = [r['b_name'] for r in cur.fetchall()]
        conn.close()
    except Exception: pass
    
    if not b_options: b_options = ["-- ไม่มีข้อมูลบอยเลอร์ --"]
    curr_b = str(row_data.get('boiler_name') or '')
    b_idx = b_options.index(curr_b) if (curr_b in b_options and e_branch_lbl == curr_branch) else 0

    with c_top2: e_boiler = st.selectbox("🔥 2. แก้ไขบอยเลอร์ *", b_options, index=b_idx)
    
    st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
    e_date = st.date_input("แก้ไข วันที่", format="DD/MM/YYYY", value=pd.to_datetime(row_data.get('record_date')), max_value=pd.to_datetime("today"))
    c1, c2, c3 = st.columns(3)
    with c1:
        st.write(" ")
        e_saw_w = st.number_input("ขี้เลื่อย (ตัน)", min_value=0.0, value=float(row_data.get('sawdust_weight') or 0.0))
        e_wood_w = st.number_input("ปีกไม้ (ตัน)", min_value=0.0, value=float(row_data.get('wood_weight') or 0.0))
        e_waste_w = st.number_input("เศษไม้เสีย (ตัน)", min_value=0.0, value=float(row_data.get('waste_wood_weight') or 0.0))
    with c2:
        st.write(" ")
        e_saw_p = st.number_input("ราคาขี้เลื่อย", min_value=0.0, value=float(row_data.get('sawdust_price') or 0.0))
        e_wood_p = st.number_input("ราคาปีกไม้", min_value=0.0, value=float(row_data.get('wood_price') or 0.0))
        e_waste_p = st.number_input("ราคาเศษไม้เสีย", min_value=0.0, value=float(row_data.get('waste_wood_price') or 0.0))
    with c3:
        st.write(" ")
        e_prod = st.number_input("ผลิตไอน้ำ (ตัน)", min_value=0.0, value=float(row_data.get('steam_production') or 1.0))
        e_hrs = st.number_input("ชม.ทำงาน", min_value=0.0, value=float(row_data.get('working_hours') or 1.0))

    btn_disabled = (e_boiler == "-- ไม่มีข้อมูลบอยเลอร์ --")
    if st.button("💾 บันทึกการแก้ไข", use_container_width=True, type="primary", disabled=btn_disabled):
        if e_hrs > 24.0:
            st.error("⚠️ ไม่สามารถบันทึกได้: ชั่วโมงทำงานต้องไม่เกิน 24 ชั่วโมงต่อวัน")
            return
        with st.spinner("กำลังบันทึกข้อมูล..."):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"""UPDATE boiler_fuel_records SET branch_id=%s, record_date=%s, boiler_name=%s, sawdust_weight=%s, wood_weight=%s, waste_wood_weight=%s, 
                                   sawdust_price=%s, wood_price=%s, waste_wood_price=%s, steam_production=%s, working_hours=%s WHERE {pk_col}=%s""",
                                (new_branch_id, e_date, e_boiler, e_saw_w, e_wood_w, e_waste_w, e_saw_p, e_wood_p, e_waste_p, e_prod, e_hrs, rec_id))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "UPDATE", "All Report: Tab 4", f"แก้ไข ID: {rec_id}")
                st.toast("✅ อัปเดตข้อมูลสำเร็จ!", icon="🔥")
                time.sleep(1.2)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("🛠️ แก้ไขข้อมูล (ยอดรวมสาขา)", width="large")
def update_record_dialog_oven(row_data, pk_col):
    rec_id = row_data[pk_col]
    st.info(f"ID: {rec_id} | สาขาเดิม: {row_data.get('branch_name')}")
    
    raw_role = st.session_state.get('role_tab') or st.session_state.get('role') or 'user'
    user_allowed_branches = st.session_state.get('allowed_branches', [str(st.session_state.branch_id)])
    if str(raw_role).strip().lower() in ["admin", "reporter"]: edit_branch_options = list(branch_dict.keys())
    else: edit_branch_options = [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
    
    curr_branch = str(row_data.get('branch_name') or '')
    def_branch_idx = edit_branch_options.index(curr_branch) if curr_branch in edit_branch_options else 0
    
    e_branch_lbl = st.selectbox("🏢 1. แก้ไขสาขา *", edit_branch_options, index=def_branch_idx, key=f"e_ov_br_{rec_id}")
    new_branch_id = branch_dict[e_branch_lbl]
    
    st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
    e_date = st.date_input("แก้ไข วันที่", format="DD/MM/YYYY", value=pd.to_datetime(row_data.get('record_date')), max_value=pd.to_datetime("today"), key=f"e_ov_dt_{rec_id}")
    
    c1, c2, c3 = st.columns(3)
    with c1: e_ov = st.number_input("แก้ไข จำนวนเตาที่อบ", min_value=0, value=int(row_data.get('oven_qty') or 0), key=f"e_ov_q_{rec_id}")
    with c2: e_wd = st.number_input("แก้ไข ไม้อบออก (ลบ.ฟ.)", min_value=0.0, value=float(row_data.get('wood_out_cubft') or 0.0), key=f"e_ov_w_{rec_id}")
    with c3: e_pr = st.number_input("แก้ไข แรงดันปลายทางเฉลี่ย", min_value=0.0, value=float(row_data.get('avg_terminal_pressure') or 0.0), key=f"e_ov_p_{rec_id}")

    if st.button("💾 บันทึกการแก้ไข", use_container_width=True, type="primary", key=f"e_ov_save_{rec_id}"):
        with st.spinner("กำลังบันทึกข้อมูล..."):
            try:
                conn = get_db_connection()
                with conn.cursor() as cur:
                    cur.execute(f"UPDATE daily_wood_oven_records SET branch_id=%s, record_date=%s, oven_qty=%s, wood_out_cubft=%s, avg_terminal_pressure=%s WHERE {pk_col}=%s",
                                (new_branch_id, e_date, e_ov, e_wd, e_pr, rec_id))
                    conn.commit()
                conn.close()
                st.cache_data.clear()
                log_activity(st.session_state.user_id, st.session_state.username, "UPDATE", "All Report: Tab 4 Oven", f"แก้ไข ID: {rec_id}")
                st.toast("✅ อัปเดตข้อมูลสำเร็จ!", icon="📁")
                time.sleep(1.2)
                st.rerun()
            except Exception as e: st.error(f"Error: {e}")

@st.dialog("📊 หน้าต่างดูสรุปรายงาน", width="large")
def show_summary_report_dialog(html_content):
    components.html(html_content, height=650, scrolling=True)

@st.dialog("📊 หน้าต่างดูสรุปรายงาน", width="large")
def show_summary_report_dialog_t4(boiler_dict):
    boilers = list(boiler_dict.keys())
    if not boilers:
        st.info("ไม่มีข้อมูลรายงาน")
        return
        
    tab_names = [b_name for b_name in boilers]
    tabs = st.tabs(tab_names)
    
    for idx, b_name in enumerate(boilers):
        with tabs[idx]:
            st.components.v1.html(boiler_dict[b_name], height=600, scrolling=True)

# ==========================================================
# 🚀 3. ฟังก์ชันกลางจัดการ Pagination
# ==========================================================
def render_custom_pagination(data_list, session_prefix, default_rows=15):
    total_records = len(data_list)
    rows_key = f"{session_prefix}_rows"
    page_key = f"{session_prefix}_page"

    if rows_key not in st.session_state: st.session_state[rows_key] = default_rows
    if page_key not in st.session_state: st.session_state[page_key] = 1

    rows_per_page = st.session_state[rows_key]
    
    if rows_per_page <= 0: total_pages = 1
    elif rows_per_page >= total_records: total_pages = 1
    else: total_pages = max(1, math.ceil(total_records / rows_per_page))

    if st.session_state[page_key] > total_pages: st.session_state[page_key] = total_pages
    if st.session_state[page_key] < 1: st.session_state[page_key] = 1

    pc1, pc_space, pc2 = st.columns([3.5, 2.5, 4.0])
    with pc1: 
        new_rows = st.number_input(f"พบข้อมูลทั้งหมด {total_records} รายการ (แสดงหน้าละ):", value=st.session_state[rows_key], step=5, key=f"inp_{rows_key}")
        if new_rows != st.session_state[rows_key]:
            st.session_state[rows_key] = new_rows
            st.session_state[page_key] = 1
            st.rerun()

    with pc2: 
        st.write("")
        dyn_key = f"sel_pg_{session_prefix}_{st.session_state[page_key]}"
        selected_page = st.selectbox("เลือกหน้า", range(1, total_pages + 1), index=st.session_state[page_key] - 1, format_func=lambda x: f"📑 หน้า {x} / {total_pages}", key=dyn_key, label_visibility="collapsed")
        if selected_page != st.session_state[page_key]:
            st.session_state[page_key] = selected_page
            st.rerun()

        st.markdown("<div style='margin-top: -10px;'></div>", unsafe_allow_html=True)
        c_prev, c_page, c_next = st.columns([1.2, 1.6, 1.2])
        with c_prev:
            if st.button("⬅️ ก่อนหน้า", use_container_width=True, disabled=(st.session_state[page_key] <= 1), key=f"btn_prev_{session_prefix}"):
                st.session_state[page_key] -= 1
                st.rerun()
        with c_page: 
            st.markdown(f"<div style='text-align:center; padding-top:8px; font-weight:bold; color:#57534E; font-size:14px;'>หน้า {st.session_state[page_key]} / {total_pages}</div>", unsafe_allow_html=True)
        with c_next:
            if st.button("ถัดไป ➡️", use_container_width=True, disabled=(st.session_state[page_key] >= total_pages), key=f"btn_next_{session_prefix}"):
                st.session_state[page_key] += 1
                st.rerun()

    if rows_per_page <= 0: return []
    elif rows_per_page >= total_records: return data_list
    else:
        start_idx = (st.session_state[page_key] - 1) * rows_per_page
        end_idx = start_idx + rows_per_page
        return data_list[start_idx:end_idx]


# ==========================================================
# 🚀 4. ฟังก์ชัน CACHE ประมวลผลข้อมูล
# ==========================================================
@st.cache_data(show_spinner=False)
def process_t1(b_id_t1, start_date_str, end_date_str, allowed_tuple, selected_branch_display):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            if b_id_t1 == "ทั้งหมด":
                sql = "SELECT t.*, b.branch_name, m.machine_name FROM machine_trans t LEFT JOIN branches b ON t.branch_id = b.id LEFT JOIN machines_set m ON t.machine_id = m.id WHERE t.record_date BETWEEN %s AND %s ORDER BY t.record_date ASC"
                cur.execute(sql, (start_date_str, end_date_str))
            elif b_id_t1 == "รวมเฉพาะที่มีสิทธิ์":
                placeholders = ', '.join(['%s'] * len(allowed_tuple))
                sql = f"SELECT t.*, b.branch_name, m.machine_name FROM machine_trans t LEFT JOIN branches b ON t.branch_id = b.id LEFT JOIN machines_set m ON t.machine_id = m.id WHERE t.branch_id IN ({placeholders}) AND t.record_date BETWEEN %s AND %s ORDER BY t.record_date ASC"
                cur.execute(sql, allowed_tuple + (start_date_str, end_date_str))
            else:
                sql = "SELECT t.*, b.branch_name, m.machine_name FROM machine_trans t LEFT JOIN branches b ON t.branch_id = b.id LEFT JOIN machines_set m ON t.machine_id = m.id WHERE t.branch_id = %s AND t.record_date BETWEEN %s AND %s ORDER BY t.record_date ASC"
                cur.execute(sql, (b_id_t1, start_date_str, end_date_str))
            raw_data = cur.fetchall()
        conn.close()
    except Exception as e:
        st.error(f"⚠️ เกิดข้อผิดพลาดในการดึงข้อมูล: {e}")
        return [], None, None

    if not raw_data: return [], None, None

    unique_m_names = set()
    for r in raw_data:
        nm = str(r.get('machine_name') or '').strip()
        if nm: unique_m_names.add(nm)
        
    legacy_mapping = {
        "โต๊ะเลื่อย": {"suffix": " (≤ 0.15%)", "hex": "E2EFDA"},
        "พัดลม": {"suffix": " (≤ 0.15%)", "hex": "F2F2F2"},
        "ตาชั่งใหญ่": {"suffix": " (≤ 0.20%)", "hex": "EDEDED"},
        "ตาชั่งเล็ก": {"suffix": " (≤ 0.20%)", "hex": "F2F2F2"},
        "รถยก": {"suffix": " (≤ 0.20%)", "hex": "FCE4D6"},
        "รถคีบ": {"suffix": " (≤16 ชม.)", "hex": "FFF2CC"},
        "อัดน้ำยา": {"suffix": " (≤ 0.20%)", "hex": "E2EFDA"},
        "เตาอบปกติ": {"suffix": " (≤ 0.20%)", "hex": "D9E1F2"},
        "cdk": {"suffix": " (≤ 0.20%)", "hex": "D9E1F2"},
        "บอยเลอร์": {"suffix": " (≤ 0.15%)", "hex": "FCE4D6"},
        "ชิปเปอร์": {"suffix": " (≤ 0.15%)", "hex": "FFF2CC"}
    }
    
    color_palette = ["D9E1F2", "E2EFDA", "FFF2CC", "FCE4D6", "F2F2F2", "EDEDED"]
    new_color_idx = 0
    machines_config = []
    display_name_map = {} 
    
    for m_name in sorted(list(unique_m_names)):
        disp_name = m_name
        assigned_hex = None
        for k, v in legacy_mapping.items():
            if k.lower() in m_name.lower():
                # 🌟 แก้บัคชื่อเบิ้ล (เช่น โต๊ะเลื่อย (<= 0.15%) (<= 0.15%))
                if v['suffix'].strip() not in m_name:
                    disp_name = f"{m_name}{v['suffix']}"
                else:
                    disp_name = m_name
                assigned_hex = v["hex"]
                break
        if not assigned_hex:
            assigned_hex = color_palette[new_color_idx % len(color_palette)]
            new_color_idx += 1
        machines_config.append({"name": disp_name, "hex": assigned_hex})
        display_name_map[m_name] = disp_name
        
    if not machines_config: return [], None, None

    start_dt = pd.to_datetime(start_date_str)
    days_in_month = calendar.monthrange(start_dt.year, start_dt.month)[1]
    
    # 🌟 ตัวแปรดักจับว่า "เครื่องจักรไหนมีหมายเหตุบ้าง"
    machine_has_remarks = {m['name']: False for m in machines_config}
    
    matrix_data = {d: {m['name']: {'qty': 0, 'work': 0.0, 'break': 0.0, 'rem': '', 'has_data': False} for m in machines_config} for d in range(1, days_in_month + 1)}
    
    for row_m in raw_data:
        d_num = pd.to_datetime(row_m['record_date']).day
        m_db_name = str(row_m.get('machine_name') or '').strip()
        if not m_db_name: continue
        mapped_name = display_name_map[m_db_name]
        w_hr = float(row_m.get('working_hours') or 0.0)
        b_hr = float(row_m.get('breakdown_hours') or 0.0)
        qty = int(row_m.get('machine_qty') or 0)
        rem = str(row_m.get('remarks') or '').strip()
        
        matrix_data[d_num][mapped_name]['qty'] += qty
        matrix_data[d_num][mapped_name]['work'] += w_hr
        matrix_data[d_num][mapped_name]['break'] += b_hr
        matrix_data[d_num][mapped_name]['has_data'] = True
        
        # เก็บหมายเหตุ ถ้ามี
        if rem and rem != '-' and rem.lower() != 'none':
            machine_has_remarks[mapped_name] = True
            if matrix_data[d_num][mapped_name]['rem']:
                matrix_data[d_num][mapped_name]['rem'] += f", {rem}"
            else:
                matrix_data[d_num][mapped_name]['rem'] = rem

    active_machines_config = []
    for m in machines_config:
        if any([matrix_data[d][m['name']]['has_data'] for d in range(1, days_in_month + 1)]):
            # 🌟 กำหนดจำนวนคอลัมน์แบบไดนามิก: มีหมายเหตุ=4, ไม่มีหมายเหตุ=3
            m['has_rem'] = machine_has_remarks[m['name']]
            m['cols'] = 4 if m['has_rem'] else 3
            active_machines_config.append(m)
            
    if not active_machines_config: active_machines_config = machines_config
        
    totals_work = {m['name']: sum([matrix_data[d][m['name']]['work'] for d in range(1, days_in_month + 1)]) for m in active_machines_config}
    totals_break = {m['name']: sum([matrix_data[d][m['name']]['break'] for d in range(1, days_in_month + 1)]) for m in active_machines_config}
    percentages = {m['name']: f"{(totals_break[m['name']] / totals_work[m['name']] * 100):.2f}%" if totals_work[m['name']] > 0 else "0.00%" for m in active_machines_config}
    
    start_dt = pd.to_datetime(start_date_str)
    month_str = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"][start_dt.month - 1]
    year_buddhist = start_dt.year + 543

    # ==========================================
    # สร้างไฟล์ Excel
    # ==========================================
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary Machine Report"
    ws.views.sheetView[0].showGridLines = True
    font_title, font_banner, font_header, font_body, font_total, font_pct = Font(name="Sarabun", size=14, bold=True), Font(name="Sarabun", size=12, bold=True), Font(name="Sarabun", size=9, bold=True), Font(name="Sarabun", size=9), Font(name="Sarabun", size=9, bold=True), Font(name="Sarabun", size=10, bold=True, color="FF0000")
    align_center, align_right, align_left = Alignment(horizontal="center", vertical="center", wrap_text=True), Alignment(horizontal="right", vertical="center"), Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    fill_banner, fill_green = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"), PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    # 🌟 คำนวณคอลัมน์รวมทั้งหมดแบบไดนามิก
    total_cols = 1 + sum([m['cols'] for m in active_machines_config])
    last_col_letter = get_column_letter(total_cols)

    ws.merge_cells(f"A1:{last_col_letter}1"); ws["A1"] = f"บริษัท วู้ดเวิร์ค จำกัด สาขา {selected_branch_display}"; ws["A1"].font, ws["A1"].alignment = font_title, align_center
    ws.merge_cells(f"A2:{last_col_letter}2"); ws["A2"] = f"สรุปชั่วโมงการทำงานของเครื่องจักร/เบรกดาวน์ ประจำเดือน....{month_str}................... {year_buddhist}"; ws["A2"].font, ws["A2"].alignment, ws["A2"].fill = font_banner, align_center, fill_banner
    ws.row_dimensions[2].height = 25
    ws.merge_cells("A3:A4"); ws["A3"] = "วันที่"; ws["A3"].font, ws["A3"].alignment, ws["A3"].fill, ws["A3"].border = font_header, align_center, fill_green, thin_border
    ws["A4"].border = thin_border

    col_idx = 2
    for m in active_machines_config:
        cols_span = m['cols']
        start_c, end_c = get_column_letter(col_idx), get_column_letter(col_idx + cols_span - 1)
        ws.merge_cells(f"{start_c}3:{end_c}3")
        ws[f"{start_c}3"] = m["name"]; ws[f"{start_c}3"].font, ws[f"{start_c}3"].alignment = font_header, align_center; ws[f"{start_c}3"].fill = PatternFill(start_color=m["hex"], end_color=m["hex"], fill_type="solid")
        for c in range(col_idx, col_idx + cols_span): ws[f"{get_column_letter(c)}3"].border = thin_border
        
        sub_headers = ["เครื่องจักร\nใช้งาน\n(เครื่อง)", "ชั่วโมง\nทำงาน\n(ชม.)", "ชั่วโมง\nเบรกดาวน์\n(ชม.)"]
        if m['has_rem']: sub_headers.append("หมายเหตุ")
        
        for i, sh in enumerate(sub_headers):
            cell_ref = f"{get_column_letter(col_idx + i)}4"; ws[cell_ref] = sh; ws[cell_ref].font, ws[cell_ref].alignment, ws[cell_ref].border = Font(name="Sarabun", size=8, bold=True), align_center, thin_border
        col_idx += cols_span

    ws.row_dimensions[3].height, ws.row_dimensions[4].height = 20, 35
    current_row = 5
    for day in range(1, days_in_month + 1):
        ws[f"A{current_row}"] = day; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_header, align_center, thin_border
        c_idx = 2
        for m in active_machines_config:
            item = matrix_data[day][m['name']]
            cell_q, cell_w, cell_b = ws[f"{get_column_letter(c_idx)}{current_row}"], ws[f"{get_column_letter(c_idx+1)}{current_row}"], ws[f"{get_column_letter(c_idx+2)}{current_row}"]
            cell_q.value = (item['qty'] if item['qty'] > 0 else "-") if item['has_data'] else ""
            cell_w.value = (item['work'] if item['work'] > 0 else "-") if item['has_data'] else ""
            cell_b.value = (item['break'] if item['break'] > 0 else "-") if item['has_data'] else ""
            
            cells_to_style = [(cell_q, align_center), (cell_w, align_right if isinstance(cell_w.value, (int, float)) else align_center), (cell_b, align_right if isinstance(cell_b.value, (int, float)) else align_center)]
            
            if m['has_rem']:
                cell_r = ws[f"{get_column_letter(c_idx+3)}{current_row}"]
                cell_r.value = item['rem'] if item['has_data'] else ""
                cells_to_style.append((cell_r, align_left))
                
            for cell, al in cells_to_style:
                cell.font, cell.alignment, cell.border = font_body, al, thin_border
                if isinstance(cell.value, (int, float)): cell.number_format = '#,##0.00'
            c_idx += m['cols']
        current_row += 1

    ws[f"A{current_row}"] = "รวม"; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].fill, ws[f"A{current_row}"].border = font_total, align_center, fill_green, thin_border
    c_idx = 2
    for m in active_machines_config:
        m_n = m['name']
        cell_q, cell_w, cell_b = ws[f"{get_column_letter(c_idx)}{current_row}"], ws[f"{get_column_letter(c_idx+1)}{current_row}"], ws[f"{get_column_letter(c_idx+2)}{current_row}"]
        cell_q.value = "-"
        cell_w.value = totals_work[m_n] if totals_work[m_n] > 0 else "-"
        cell_b.value = totals_break[m_n] if totals_break[m_n] > 0 else "-"
        
        cells_to_style = [(cell_q, align_center), (cell_w, align_right if totals_work[m_n] > 0 else align_center), (cell_b, align_right if totals_break[m_n] > 0 else align_center)]
        
        if m['has_rem']:
            cell_r = ws[f"{get_column_letter(c_idx+3)}{current_row}"]
            cell_r.value = ""
            cells_to_style.append((cell_r, align_center))

        for cell, al in cells_to_style:
            cell.font, cell.alignment, cell.fill, cell.border = font_total, al, fill_green, thin_border
            if isinstance(cell.value, (int, float)): cell.number_format = '#,##0.00'
        c_idx += m['cols']

    current_row += 1
    ws[f"A{current_row}"] = "คิดเป็น%"; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_pct, align_center, thin_border
    c_idx = 2
    for m in active_machines_config:
        cols_span = m['cols']
        start_c, end_c = get_column_letter(c_idx), get_column_letter(c_idx + cols_span - 1)
        ws.merge_cells(f"{start_c}{current_row}:{end_c}{current_row}")
        ws[f"{start_c}{current_row}"].value = percentages[m['name']]; ws[f"{start_c}{current_row}"].font, ws[f"{start_c}{current_row}"].alignment = font_pct, align_center
        for c in range(c_idx, c_idx + cols_span): ws[f"{get_column_letter(c)}{current_row}"].border = thin_border
        c_idx += cols_span

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_data = excel_buffer.getvalue()

    # ==========================================
    # สร้างตาราง HTML (Pop-up สรุปรายงาน & Print)
    # ==========================================
    MACHINES_PER_PAGE = 7 
    machine_chunks = [active_machines_config[i:i + MACHINES_PER_PAGE] for i in range(0, len(active_machines_config), MACHINES_PER_PAGE)]
    all_tables_html = ""
    
    for chunk_idx, chunk_machines in enumerate(machine_chunks):
        # 🌟 กำหนด colspan ให้ตรงตามสถานะของเครื่องจักร
        header_row1 = "".join([f'<th colspan="{m["cols"]}" style="background-color:#{m["hex"]};border:1px solid #000;padding:3px;font-size:10px;text-align:center;">{m["name"]}</th>' for m in chunk_machines])
        
        header_row2 = ""
        for m in chunk_machines:
            header_row2 += '<th style="border:1px solid #000;padding:2px;font-size:8px;width:28px;">เครื่องจักร<br>ใช้งาน</th><th style="border:1px solid #000;padding:2px;font-size:8px;width:32px;">ชั่วโมง<br>ทำงาน</th><th style="border:1px solid #000;padding:2px;font-size:8px;width:32px;">ชั่วโมง<br>เบรกดาวน์</th>'
            if m['has_rem']:
                header_row2 += '<th style="border:1px solid #000;padding:2px;font-size:8px;width:45px;">หมายเหตุ</th>'

        body_rows = ""
        for day in range(1, days_in_month + 1):
            body_rows += f'<tr><td style="border:1px solid #000;padding:2px;text-align:center;font-size:9px;font-weight:bold;">{day}</td>'
            for m in chunk_machines:
                item = matrix_data[day][m['name']]
                if item['has_data']:
                    q_val = str(item['qty']) if item['qty'] > 0 else "-"
                    w_val = f"{item['work']:,.2f}" if item['work'] > 0 else "-"
                    b_val = f"{item['break']:,.2f}" if item['break'] > 0 else "-"
                    r_val = item['rem']
                else: q_val, w_val, b_val, r_val = "", "", "", ""
                
                body_rows += f'<td style="border:1px solid #000;padding:2px;text-align:center;font-size:9px;">{q_val}</td><td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{w_val}</td><td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{b_val}</td>'
                if m['has_rem']:
                    body_rows += f'<td style="border:1px solid #000;padding:2px;text-align:left;font-size:8px;color:#475569;">{r_val}</td>'
            body_rows += '</tr>'
            
        total_row_html = '<tr><td style="border:1px solid #000;padding:3px;text-align:center;font-size:10px;font-weight:bold;background-color:#E2EFDA;">รวม</td>'
        for m in chunk_machines:
            m_n = m['name']
            tw_s = f"{totals_work[m_n]:,.2f}" if totals_work[m_n] > 0 else '-'
            tb_s = f"{totals_break[m_n]:,.2f}" if totals_break[m_n] > 0 else '-'
            total_row_html += f'<td style="border:1px solid #000;padding:3px;text-align:center;font-size:9px;font-weight:bold;background-color:#E2EFDA;">-</td><td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;background-color:#E2EFDA;">{tw_s}</td><td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;background-color:#E2EFDA;">{tb_s}</td>'
            if m['has_rem']:
                total_row_html += '<td style="border:1px solid #000;padding:3px;text-align:center;font-size:9px;font-weight:bold;background-color:#E2EFDA;">-</td>'
        total_row_html += '</tr>'
        
        pct_row_html = '<tr><td style="border:1px solid #000;padding:3px;text-align:center;font-size:10px;font-weight:bold;color:red;">คิดเป็น%</td>'
        for m in chunk_machines: 
            pct_row_html += f'<td colspan="{m["cols"]}" style="border:1px solid #000;padding:3px;text-align:center;font-size:10px;font-weight:bold;color:red;">{percentages[m["name"]]}</td>'
        pct_row_html += '</tr>'
        
        page_info = f"<div style='text-align: right; font-size: 10px; margin-bottom: 2px;'>แผ่นที่ {chunk_idx + 1}/{len(machine_chunks)}</div>" if len(machine_chunks) > 1 else ""
        single_table = f"""<div class="header-title">บริษัท วู้ดเวิร์ค จำกัด สาขา {selected_branch_display}</div><div class="banner">สรุปชั่วโมงการทำงานของเครื่องจักร/เบรกดาวน์ ประจำเดือน....{month_str}................... {year_buddhist}</div>{page_info}<table><thead><tr><th rowspan="2" style="background-color:#E2EFDA;border:1px solid #000;padding:3px;font-size:10px;width:35px;color:#000000 !important;">วันที่</th>{header_row1}</tr><tr>{header_row2}</tr></thead><tbody>{body_rows}{total_row_html}{pct_row_html}</tbody></table>"""
        all_tables_html += single_table
        if chunk_idx < len(machine_chunks) - 1: all_tables_html += "<div style='page-break-after: always; margin-bottom: 30px;'></div>"

    table_full_html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>Print</title><style>@page {{ size: A4 landscape; margin: 5mm; }} body {{ font-family: 'Sarabun', Tahoma, sans-serif; margin: 0; padding: 5px; background-color: #FFFFFF; color: #000000; -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }} .header-title {{ text-align: center; font-size: 16px; font-weight: bold; margin-bottom: 4px; }} .banner {{ background-color: #F8CBAD; text-align: center; font-size: 14px; font-weight: bold; padding: 5px; border: 1px solid #000; margin-bottom: 4px; }} table {{ width: 100%; border-collapse: collapse; }} th, td {{ font-family: 'Sarabun', Tahoma, sans-serif; color: #000000 !important; }}</style></head><body>{all_tables_html}</body></html>"""
    
    return raw_data, excel_data, json.dumps(table_full_html)

@st.cache_data(show_spinner=False)
def process_t2(b_id_t2, start_date_str, end_date_str, allowed_tuple, selected_branch_display):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            if b_id_t2 == "ทั้งหมด":
                sql_eng = "SELECT e.id AS engine_pk_id, CONVERT(e.machine_name USING utf8mb4) AS engine_code, CONVERT(et.type_name USING utf8mb4) AS type_name, CONVERT(b.branch_name USING utf8mb4) AS branch_name, e.branch_id FROM machines_set e LEFT JOIN engine_types et ON e.engine_type_id = et.id LEFT JOIN branches b ON e.branch_id = b.id WHERE e.machine_type = 'engine' AND e.is_active = 1 ORDER BY b.id ASC, e.machine_name ASC"
                cur.execute(sql_eng)
                engines_list = cur.fetchall()
                sql_trans = "SELECT r.*, b.branch_name FROM fuel_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql_trans, (start_date_str, end_date_str))
                raw_data = cur.fetchall()
            elif b_id_t2 == "รวมเฉพาะที่มีสิทธิ์":
                placeholders = ', '.join(['%s'] * len(allowed_tuple))
                sql_eng = f"SELECT e.id AS engine_pk_id, CONVERT(e.machine_name USING utf8mb4) AS engine_code, CONVERT(et.type_name USING utf8mb4) AS type_name, CONVERT(b.branch_name USING utf8mb4) AS branch_name, e.branch_id FROM machines_set e LEFT JOIN engine_types et ON e.engine_type_id = et.id LEFT JOIN branches b ON e.branch_id = b.id WHERE e.branch_id IN ({placeholders}) AND e.machine_type = 'engine' AND e.is_active = 1 ORDER BY e.machine_name ASC"
                cur.execute(sql_eng, allowed_tuple)
                engines_list = cur.fetchall()
                sql_trans = f"SELECT r.*, b.branch_name FROM fuel_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.branch_id IN ({placeholders}) AND r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql_trans, allowed_tuple + (start_date_str, end_date_str))
                raw_data = cur.fetchall()
            else:
                sql_eng = "SELECT e.id AS engine_pk_id, CONVERT(e.machine_name USING utf8mb4) AS engine_code, CONVERT(et.type_name USING utf8mb4) AS type_name, CONVERT(b.branch_name USING utf8mb4) AS branch_name, e.branch_id FROM machines_set e LEFT JOIN engine_types et ON e.engine_type_id = et.id LEFT JOIN branches b ON e.branch_id = b.id WHERE e.branch_id = %s AND e.machine_type = 'engine' AND e.is_active = 1 ORDER BY e.id ASC"
                cur.execute(sql_eng, (b_id_t2,))
                engines_list = cur.fetchall()
                sql_trans = "SELECT r.*, b.branch_name FROM fuel_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.branch_id = %s AND r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql_trans, (b_id_t2, start_date_str, end_date_str))
                raw_data = cur.fetchall()
        conn.close()
    except Exception as e:
        st.error(f"⚠️ เกิดข้อผิดพลาดในการดึงข้อมูล: {e}")
        return [], None, None, [], {}

    engine_options_t2, engine_details_t2 = {}, {}
    if engines_list:
        for eng in engines_list:
            code = str(eng.get('engine_code') or '').strip()
            t_name = str(eng.get('type_name') or 'ไม่ระบุประเภท').strip()
            b_name = str(eng.get('branch_name') or '').strip()
            try: branch_code = b_name.split("สาขา ")[1].split()[0]
            except: branch_code = b_name or "N/A"
            lbl = f"{code} (ประเภท: {t_name}) [{branch_code}]"
            engine_options_t2[lbl] = eng['engine_pk_id']
            engine_details_t2[lbl] = {'code': code, 'type': t_name}
    engine_lbl_list = list(engine_options_t2.keys())

    if not raw_data: return [], None, None, engine_lbl_list, engine_details_t2

    real_type_map = {}
    for r in raw_data:
        code, b_name, t_name = str(r.get('engine_code') or '').strip(), str(r.get('branch_name') or '').strip(), str(r.get('type_name') or '').strip()
        if b_id_t2 in ["ทั้งหมด", "รวมเฉพาะที่มีสิทธิ์"] and b_name:
            try: b_short = b_name.split("สาขา ")[1].split()[0]
            except: b_short = b_name
            k = f"{code}_{b_short}".lower()
        else: k = code.lower()
        if t_name: real_type_map[k] = t_name

    seen_keys, all_engines_config = set(), []
    if engines_list:
        for eng in engines_list:
            code, b_name = str(eng.get('engine_code') or '').strip(), str(eng.get('branch_name') or '').strip()
            if b_id_t2 in ["ทั้งหมด", "รวมเฉพาะที่มีสิทธิ์"] and b_name:
                try: b_short = b_name.split("สาขา ")[1].split()[0]
                except: b_short = b_name
                eng_key, title_name = f"{code}_{b_short}", f"{code} [{b_short}]"
            else: eng_key, title_name = code, code
            if not code or eng_key.lower() in seen_keys: continue
            seen_keys.add(eng_key.lower())
            type_nm = real_type_map.get(eng_key.lower()) or str(eng.get('type_name') or 'TOYOTA')
            hex_color = "FFC000" if "tck" in type_nm.lower() else "FF0000"
            all_engines_config.append({"key": eng_key, "code": code, "title": title_name, "brand": type_nm, "branch": b_name, "hex": hex_color})

    start_dt = pd.to_datetime(start_date_str)
    days_in_month = calendar.monthrange(start_dt.year, start_dt.month)[1]
    
    # 🌟 เพิ่ม 'rem' สำหรับเก็บหมายเหตุของรถยนต์
    matrix_data = {d: {e['key']: {'liters': 0.0, 'hours': 0.0, 'rem': '', 'has_data': False} for e in all_engines_config} for d in range(1, days_in_month + 1)}
    
    for r in raw_data:
        d_num = pd.to_datetime(r['record_date']).day
        e_code, b_name = str(r.get('engine_code') or '').strip(), str(r.get('branch_name') or '').strip()
        rem = str(r.get('remark') or '').strip()
        
        if b_id_t2 in ["ทั้งหมด", "รวมเฉพาะที่มีสิทธิ์"] and b_name:
            try: b_short = b_name.split("สาขา ")[1].split()[0]
            except: b_short = b_name
            match_key = f"{e_code}_{b_short}"
        else: match_key = e_code
            
        for e in all_engines_config:
            if e['key'].lower() == match_key.lower():
                matrix_data[d_num][e['key']]['liters'] += float(r.get('fuel_liters') or 0.0)
                matrix_data[d_num][e['key']]['hours'] += float(r.get('working_hours') or 0.0)
                matrix_data[d_num][e['key']]['has_data'] = True
                
                # 🌟 เก็บหมายเหตุ ถ้ามี
                if rem and rem != '-' and rem.lower() != 'none':
                    if matrix_data[d_num][e['key']]['rem']:
                        matrix_data[d_num][e['key']]['rem'] += f", {rem}"
                    else:
                        matrix_data[d_num][e['key']]['rem'] = rem
                break

    engines_config = []
    for e in all_engines_config:
        if any([matrix_data[d][e['key']]['has_data'] for d in range(1, days_in_month + 1)]):
            # 🌟 กำหนดจำนวนคอลัมน์แบบไดนามิก: มีหมายเหตุ=4, ไม่มีหมายเหตุ=3
            e['has_rem'] = any([matrix_data[d][e['key']]['rem'] for d in range(1, days_in_month + 1)])
            e['cols'] = 4 if e['has_rem'] else 3
            engines_config.append(e)
            
    if not engines_config: 
        for e in all_engines_config:
            e['has_rem'] = False
            e['cols'] = 3
        engines_config = all_engines_config

    start_dt = pd.to_datetime(start_date_str)
    days_in_current_month = pd.Period(start_dt, freq='M').days_in_month
    toyo_keys = [e['key'] for e in engines_config if "tck" not in e['brand'].lower()]
    tck_keys = [e['key'] for e in engines_config if "tck" in e['brand'].lower()]

    totals_hours, totals_liters, avg_l_hr, std_target_hours = {}, {}, {}, {}
    for e in engines_config:
        c_key = e['key']
        tot_h = sum([matrix_data[d][c_key]['hours'] for d in range(1, days_in_month + 1)])
        tot_l = sum([matrix_data[d][c_key]['liters'] for d in range(1, days_in_month + 1)])
        totals_hours[c_key], totals_liters[c_key] = tot_h, tot_l
        avg_l_hr[c_key] = f"{(tot_l / tot_h):.2f}" if tot_h > 0 else "0.00"
        is_backup_car = "สำรอง" in e['title'] or "สำรอง" in e['code'] or "สำรอง" in e.get('brand', '')
        std_target_hours[c_key] = (4 if is_backup_car else 20) * days_in_current_month

    summary_matrix = {}
    for d in range(1, days_in_month + 1):
        summary_matrix[d] = {
            'toyo': {'hours': sum([matrix_data[d][k]['hours'] for k in toyo_keys]), 'liters': sum([matrix_data[d][k]['liters'] for k in toyo_keys])},
            'tck': {'hours': sum([matrix_data[d][k]['hours'] for k in tck_keys]), 'liters': sum([matrix_data[d][k]['liters'] for k in tck_keys])}
        }

    tot_toyo_hours = sum([summary_matrix[d]['toyo']['hours'] for d in range(1, days_in_month + 1)])
    tot_toyo_liters = sum([summary_matrix[d]['toyo']['liters'] for d in range(1, days_in_month + 1)])
    avg_toyo_rate = f"{(tot_toyo_liters / tot_toyo_hours):.2f}" if tot_toyo_hours > 0 else "0.00"

    tot_tck_hours = sum([summary_matrix[d]['tck']['hours'] for d in range(1, days_in_month + 1)])
    tot_tck_liters = sum([summary_matrix[d]['tck']['liters'] for d in range(1, days_in_month + 1)])
    avg_tck_rate = f"{(tot_tck_liters / tot_tck_hours):.2f}" if tot_tck_hours > 0 else "0.00"

    tot_target_toyo = sum([std_target_hours[k] for k in toyo_keys])
    tot_target_tck = sum([std_target_hours[k] for k in tck_keys])

    month_str = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"][start_dt.month - 1]
    year_buddhist = start_dt.year + 543

    # ==========================================
    # สร้างไฟล์ Excel
    # ==========================================
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuel Report Matrix"
    ws.views.sheetView[0].showGridLines = True
    font_title, font_header, font_header_white, font_body, font_total, font_red = Font(name="Sarabun", size=14, bold=True), Font(name="Sarabun", size=9, bold=True), Font(name="Sarabun", size=9, bold=True, color="FFFFFF"), Font(name="Sarabun", size=9), Font(name="Sarabun", size=9, bold=True), Font(name="Sarabun", size=9, bold=True, color="FF0000")
    align_center, align_right, align_left = Alignment(horizontal="center", vertical="center", wrap_text=True), Alignment(horizontal="right", vertical="center"), Alignment(horizontal="left", vertical="center")
    thin_border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    fill_magenta, fill_green_sum, fill_blue_sum, fill_pink_sum = PatternFill(start_color="D90082", end_color="D90082", fill_type="solid"), PatternFill(start_color="A9D08E", end_color="A9D08E", fill_type="solid"), PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"), PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    # 🌟 คำนวณคอลัมน์รวมทั้งหมดแบบไดนามิก (บวกอีก 6 ช่องคือของกลุ่ม Toyo และ TCK)
    last_col_idx = 1 + sum([e['cols'] for e in engines_config]) + 6
    last_col_letter = get_column_letter(last_col_idx)

    ws.merge_cells(f"A1:{last_col_letter}1"); ws["A1"] = f"บริษัท วู้ดเวิร์ค จำกัด สาขา {selected_branch_display}"; ws["A1"].font, ws["A1"].alignment = font_title, align_center
    ws.merge_cells(f"A2:{last_col_letter}2"); ws["A2"] = f"รายงานการใช้เชื้อเพลิงของรถ ประจำเดือน {month_str} {year_buddhist}"; ws["A2"].font, ws["A2"].alignment = font_title, align_center
    ws.merge_cells("A3:A5"); ws["A3"] = "วันที่"; ws["A3"].font, ws["A3"].alignment, ws["A3"].border = font_header, align_center, thin_border
    ws["A4"].border, ws["A5"].border = thin_border, thin_border

    col_idx = 2
    for e in engines_config:
        cols_span = e['cols']
        start_c, end_c = get_column_letter(col_idx), get_column_letter(col_idx + cols_span - 1)
        ws.merge_cells(f"{start_c}3:{end_c}3"); ws[f"{start_c}3"] = e["title"]; ws[f"{start_c}3"].font, ws[f"{start_c}3"].alignment = font_header, align_center
        ws.merge_cells(f"{start_c}4:{end_c}4"); ws[f"{start_c}4"] = e["brand"]; ws[f"{start_c}4"].font = font_header_white if e["hex"] == "FF0000" else font_header; ws[f"{start_c}4"].alignment = align_center
        fill_color = PatternFill(start_color=e["hex"], end_color=e["hex"], fill_type="solid"); ws[f"{start_c}4"].fill = fill_color
        for c in range(col_idx, col_idx + cols_span): 
            ws[f"{get_column_letter(c)}3"].border = thin_border; ws[f"{get_column_letter(c)}4"].border = thin_border; ws[f"{get_column_letter(c)}4"].fill = fill_color
            
        sub_hdrs = ["ชม.", "ลิตร", "ล/ชม."]
        if e['has_rem']: sub_hdrs.append("หมายเหตุ")
        
        for i, sh in enumerate(sub_hdrs): 
            cell_ref = f"{get_column_letter(col_idx + i)}5"; ws[cell_ref] = sh; ws[cell_ref].font, ws[cell_ref].alignment, ws[cell_ref].border = font_header, align_center, thin_border
        col_idx += cols_span

    for grp_title, grp_bg, is_white in [("toyo", "FF0000", True), ("TCK", "FFC000", False)]:
        start_c, end_c = get_column_letter(col_idx), get_column_letter(col_idx + 2)
        ws.merge_cells(f"{start_c}3:{end_c}4"); ws[f"{start_c}3"] = grp_title; ws[f"{start_c}3"].font = font_header_white if is_white else font_header; ws[f"{start_c}3"].alignment = align_center; ws[f"{start_c}3"].fill = PatternFill(start_color=grp_bg, end_color=grp_bg, fill_type="solid")
        for c in range(col_idx, col_idx + 3):
            for r_h in [3, 4]: ws[f"{get_column_letter(c)}{r_h}"].border = thin_border; ws[f"{get_column_letter(c)}{r_h}"].fill = PatternFill(start_color=grp_bg, end_color=grp_bg, fill_type="solid")
        for i, sh in enumerate(["ชม.", "ลิตร", "ล/ชม."]): cell_ref = f"{get_column_letter(col_idx + i)}5"; ws[cell_ref] = sh; ws[cell_ref].font, ws[cell_ref].alignment, ws[cell_ref].border = font_header, align_center, thin_border
        col_idx += 3

    ws.row_dimensions[3].height, ws.row_dimensions[4].height, ws.row_dimensions[5].height = 22, 22, 20

    current_row = 6
    for day in range(1, days_in_month + 1):
        ws[f"A{current_row}"] = day; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_header, align_center, thin_border
        c_idx = 2
        for e in engines_config:
            item = matrix_data[day][e['key']]
            cell_h, cell_l, cell_r = ws[f"{get_column_letter(c_idx)}{current_row}"], ws[f"{get_column_letter(c_idx+1)}{current_row}"], ws[f"{get_column_letter(c_idx+2)}{current_row}"]
            cell_h.value = (item['hours'] if item['hours'] > 0 else "-") if item['has_data'] else ""
            cell_l.value = (item['liters'] if item['liters'] > 0 else "-") if item['has_data'] else ""
            cell_r.value = (round(item['liters'] / item['hours'], 2) if item['hours'] > 0 else "-") if item['has_data'] else ""
            
            for cell in [cell_h, cell_l, cell_r]:
                cell.font, cell.alignment, cell.border = font_body, (align_right if isinstance(cell.value, (int, float)) else align_center), thin_border
                if isinstance(cell.value, (int, float)): cell.number_format = '#,##0.00'
                
            if e['has_rem']:
                cell_rem = ws[f"{get_column_letter(c_idx+3)}{current_row}"]
                cell_rem.value = item['rem'] if item['has_data'] else ""
                cell_rem.font, cell_rem.alignment, cell_rem.border = font_body, align_left, thin_border
            
            c_idx += e['cols']

        for grp_k in ['toyo', 'tck']:
            grp_item = summary_matrix[day][grp_k]
            ws[f"{get_column_letter(c_idx)}{current_row}"] = grp_item['hours'] if grp_item['hours'] > 0 else "-"
            ws[f"{get_column_letter(c_idx+1)}{current_row}"] = grp_item['liters'] if grp_item['liters'] > 0 else "-"
            ws[f"{get_column_letter(c_idx+2)}{current_row}"] = round(grp_item['liters']/grp_item['hours'], 2) if grp_item['hours'] > 0 else "-"
            for offset in range(3):
                cell = ws[f"{get_column_letter(c_idx+offset)}{current_row}"]
                cell.font, cell.alignment, cell.border = font_body, (align_right if isinstance(cell.value, (int, float)) else align_center), thin_border
                if isinstance(cell.value, (int, float)): cell.number_format = '#,##0.00'
            c_idx += 3
        current_row += 1

    def add_summary_row(row_label, target_data, toyo_data, tck_data, bg_color=None, font_style=font_total, align_style=align_center):
        nonlocal current_row
        ws[f"A{current_row}"] = row_label
        ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_total, align_center, thin_border
        if bg_color: ws[f"A{current_row}"].fill = bg_color
        
        c_i = 2
        for e in engines_config:
            cols_span = e['cols']
            val = target_data.get(e['key'], "-")
            ws[f"{get_column_letter(c_i)}{current_row}"] = val
            cell = ws[f"{get_column_letter(c_i)}{current_row}"]
            cell.font, cell.alignment, cell.border = (font_header_white if font_style == font_header_white else font_style), align_style, thin_border
            if bg_color: cell.fill = bg_color
            if isinstance(val, (int, float)) and font_style != font_header_white: cell.number_format = '#,##0.00'
            
            if bg_color and font_style != font_total:
                ws.merge_cells(f"{get_column_letter(c_i)}{current_row}:{get_column_letter(c_i+cols_span-1)}{current_row}")
                for off in range(1, cols_span):
                    cell_off = ws[f"{get_column_letter(c_i+off)}{current_row}"]
                    cell_off.border = thin_border
                    cell_off.fill = bg_color
            else:
                for off in range(1, cols_span):
                    ws[f"{get_column_letter(c_i+off)}{current_row}"] = val if isinstance(val, str) else ""
                    ws[f"{get_column_letter(c_i+off)}{current_row}"].border = thin_border
            c_i += cols_span
            
        for sum_val in [toyo_data, tck_data]:
            cols_span = 3
            ws[f"{get_column_letter(c_i)}{current_row}"] = sum_val
            cell = ws[f"{get_column_letter(c_i)}{current_row}"]
            cell.font, cell.alignment, cell.border = (font_header_white if font_style == font_header_white else font_style), align_style, thin_border
            if bg_color: cell.fill = bg_color
            
            if bg_color and font_style != font_total:
                ws.merge_cells(f"{get_column_letter(c_i)}{current_row}:{get_column_letter(c_i+2)}{current_row}")
                for off in range(1, 3):
                    cell_off = ws[f"{get_column_letter(c_i+off)}{current_row}"]
                    cell_off.border = thin_border
                    cell_off.fill = bg_color
            else:
                for off in range(1, 3):
                    ws[f"{get_column_letter(c_i+off)}{current_row}"] = ""
                    ws[f"{get_column_letter(c_i+off)}{current_row}"].border = thin_border
            c_i += 3
        current_row += 1

    add_summary_row(f"ชม.รวม1-{days_in_current_month}", std_target_hours, tot_target_toyo, tot_target_tck, fill_magenta, font_header_white)
    
    ws[f"A{current_row}"] = f"ชม.ใช้จริง 1-{days_in_current_month}"; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_total, align_center, thin_border
    c_idx = 2
    for e in engines_config:
        c_key = e['key']
        ws[f"{get_column_letter(c_idx)}{current_row}"] = round(totals_hours[c_key], 0) if totals_hours[c_key]>0 else "-"
        ws[f"{get_column_letter(c_idx+1)}{current_row}"] = round(totals_liters[c_key], 0) if totals_liters[c_key]>0 else "-"
        ws[f"{get_column_letter(c_idx+2)}{current_row}"] = avg_l_hr[c_key]
        if e['has_rem']:
            ws[f"{get_column_letter(c_idx+3)}{current_row}"] = ""
        for offset in range(e['cols']): 
            ws[f"{get_column_letter(c_idx+offset)}{current_row}"].font, ws[f"{get_column_letter(c_idx+offset)}{current_row}"].alignment, ws[f"{get_column_letter(c_idx+offset)}{current_row}"].border = font_total, (align_right if offset < 3 else align_center), thin_border
        c_idx += e['cols']
        
    for th, tl, tr in [(tot_toyo_hours, tot_toyo_liters, avg_toyo_rate), (tot_tck_hours, tot_tck_liters, avg_tck_rate)]:
        ws[f"{get_column_letter(c_idx)}{current_row}"] = round(th, 0)
        ws[f"{get_column_letter(c_idx+1)}{current_row}"] = round(tl, 0)
        ws[f"{get_column_letter(c_idx+2)}{current_row}"] = tr
        for offset in range(3): ws[f"{get_column_letter(c_idx+offset)}{current_row}"].font, ws[f"{get_column_letter(c_idx+offset)}{current_row}"].alignment, ws[f"{get_column_letter(c_idx+offset)}{current_row}"].border = font_total, align_right, thin_border
        c_idx += 3
    current_row += 1

    ws[f"A{current_row}"] = "ส่วนต่าง"; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_total, align_center, thin_border
    c_idx = 2
    for e in engines_config:
        diff_val = std_target_hours[e['key']] - round(totals_hours[e['key']], 0)
        ws[f"{get_column_letter(c_idx)}{current_row}"] = int(diff_val); ws[f"{get_column_letter(c_idx)}{current_row}"].font, ws[f"{get_column_letter(c_idx)}{current_row}"].alignment, ws[f"{get_column_letter(c_idx)}{current_row}"].border = font_red if diff_val < 0 else font_total, align_center, thin_border
        for offset in range(1, e['cols']): ws[f"{get_column_letter(c_idx+offset)}{current_row}"].border = thin_border
        c_idx += e['cols']
        
    for diff_val in [tot_target_toyo - round(tot_toyo_hours, 0), tot_target_tck - round(tot_tck_hours, 0)]:
        ws[f"{get_column_letter(c_idx)}{current_row}"] = int(diff_val); ws[f"{get_column_letter(c_idx)}{current_row}"].font, ws[f"{get_column_letter(c_idx)}{current_row}"].alignment, ws[f"{get_column_letter(c_idx)}{current_row}"].border = font_red if diff_val < 0 else font_total, align_center, thin_border
        for offset in range(1, 3): ws[f"{get_column_letter(c_idx+offset)}{current_row}"].border = thin_border
        c_idx += 3
    current_row += 1

    add_summary_row("ชม.รวมทั้งเดือน", std_target_hours, tot_target_toyo, tot_target_tck, fill_blue_sum, font_total, align_center)
    diff_dict = {e['key']: std_target_hours[e['key']] - round(totals_hours[e['key']], 0) for e in engines_config}
    add_summary_row("ชม.คงเหลือทั้งเดือน", diff_dict, tot_target_toyo - round(tot_toyo_hours, 0), tot_target_tck - round(tot_tck_hours, 0), fill_pink_sum, font_red, align_center)

    ws.column_dimensions['A'].width = 16
    for c in range(2, last_col_idx + 1): ws.column_dimensions[get_column_letter(c)].width = 10

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    
    # ==========================================
    # สร้างตาราง HTML (Pop-up สรุปรายงาน & Print)
    # ==========================================
    # 🌟 ลดจำนวนรถต่อหน้าเหลือ 6 เพื่อเผื่อคอลัมน์ให้หน้าสุดท้าย
    ENGINES_PER_PAGE = 6 
    engine_chunks = [engines_config[i:i + ENGINES_PER_PAGE] for i in range(0, len(engines_config), ENGINES_PER_PAGE)]
    all_tables_html = ""
    for chunk_idx, chunk_engines in enumerate(engine_chunks):
        is_last_chunk = (chunk_idx == len(engine_chunks) - 1) 
        
        header_row1 = "".join([f'<th colspan="{e["cols"]}" style="border:1px solid #000;padding:3px;font-size:10px;text-align:center;">{e["title"]}</th>' for e in chunk_engines])
        if is_last_chunk: header_row1 += '<th colspan="3" style="background-color:#FF0000;color:#FFF;border:1px solid #000;padding:3px;font-size:10px;text-align:center;" rowspan="2">toyo</th><th colspan="3" style="background-color:#FFC000;border:1px solid #000;padding:3px;font-size:10px;text-align:center;" rowspan="2">TCK</th>'
        
        header_row2 = "".join([f'<th colspan="{e["cols"]}" style="background-color:#{e["hex"]};color:{"#FFF" if e["hex"]=="FF0000" else "#000"};border:1px solid #000;padding:3px;font-size:10px;text-align:center;">{e["brand"]}</th>' for e in chunk_engines])
        
        header_row3 = ""
        for e in chunk_engines:
            header_row3 += '<th style="border:1px solid #000;padding:2px;font-size:8px;width:30px;">ชม.</th><th style="border:1px solid #000;padding:2px;font-size:8px;width:30px;">ลิตร</th><th style="border:1px solid #000;padding:2px;font-size:8px;width:30px;">ล/ชม.</th>'
            if e['has_rem']:
                header_row3 += '<th style="border:1px solid #000;padding:2px;font-size:8px;width:45px;">หมายเหตุ</th>'
                
        if is_last_chunk:
            for _ in range(2):
                header_row3 += '<th style="border:1px solid #000;padding:2px;font-size:8px;width:30px;">ชม.</th><th style="border:1px solid #000;padding:2px;font-size:8px;width:30px;">ลิตร</th><th style="border:1px solid #000;padding:2px;font-size:8px;width:30px;">ล/ชม.</th>'
                
        body_rows = ""
        for day in range(1, days_in_month + 1):
            body_rows += f'<tr><td style="border:1px solid #000;padding:2px;text-align:center;font-size:9px;font-weight:bold;">{day}</td>'
            for e in chunk_engines:
                item = matrix_data[day][e['key']]
                h_val, l_val, r_val = (f"{item['hours']:,.2f}", f"{item['liters']:,.2f}", f"{(item['liters']/item['hours']):,.2f}") if item['has_data'] and item['hours']>0 else ("-", "-", "-")
                rem_val = item['rem'] if item['has_data'] else ""
                
                body_rows += f'<td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{h_val}</td><td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{l_val}</td><td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{r_val}</td>'
                if e['has_rem']:
                    body_rows += f'<td style="border:1px solid #000;padding:2px;text-align:left;font-size:8px;color:#475569;">{rem_val}</td>'
                    
            if is_last_chunk:
                for grp_k in ['toyo', 'tck']:
                    grp_item = summary_matrix[day][grp_k]
                    th_s, tl_s, tr_s = (f"{grp_item['hours']:,.2f}", f"{grp_item['liters']:,.2f}", f"{(grp_item['liters']/grp_item['hours']):,.2f}") if grp_item['hours']>0 else ("-", "-", "-")
                    body_rows += f'<td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{th_s}</td><td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{tl_s}</td><td style="border:1px solid #000;padding:2px;text-align:right;font-size:9px;">{tr_s}</td>'
            body_rows += '</tr>'
            
        total_row_html = '<tr><td style="border:1px solid #000;padding:3px;text-align:center;font-size:10px;font-weight:bold;">รวม</td>'
        for e in chunk_engines:
            c_key = e['key']
            t_h = f"{totals_hours[c_key]:,.2f}" if totals_hours[c_key] > 0 else "-"
            t_l = f"{totals_liters[c_key]:,.2f}" if totals_liters[c_key] > 0 else "-"
            t_r = avg_l_hr[c_key]
            total_row_html += f'<td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;">{t_h}</td><td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;">{t_l}</td><td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;">{t_r}</td>'
            if e['has_rem']:
                total_row_html += f'<td style="border:1px solid #000;padding:3px;text-align:center;font-size:9px;font-weight:bold;">-</td>'
                
        if is_last_chunk:
            for th, tl, tr in [(tot_toyo_hours, tot_toyo_liters, avg_toyo_rate), (tot_tck_hours, tot_tck_liters, avg_tck_rate)]:
                th_s = f"{th:,.2f}" if th > 0 else "-"
                tl_s = f"{tl:,.2f}" if tl > 0 else "-"
                total_row_html += f'<td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;">{th_s}</td><td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;">{tl_s}</td><td style="border:1px solid #000;padding:3px;text-align:right;font-size:9px;font-weight:bold;">{tr}</td>'
        total_row_html += '</tr>'
        
        page_info = f"<div style='text-align: right; font-size: 10px; margin-bottom: 2px;'>แผ่นที่ {chunk_idx + 1}/{len(engine_chunks)}</div>" if len(engine_chunks) > 1 else ""
        single_table = f"<div class='header-title'>บริษัท วู้ดเวิร์ค จำกัด สาขา {selected_branch_display}</div><div class='banner'>รายงานการใช้เชื้อเพลิงของรถ ประจำเดือน {month_str} {year_buddhist}</div>{page_info}<table><thead><tr><th rowspan='3' style='border:1px solid #000;padding:3px;font-size:10px;width:35px;'>วันที่</th>{header_row1}</tr><tr>{header_row2}</tr><tr>{header_row3}</tr></thead><tbody>{body_rows}{total_row_html}</tbody></table>"
        all_tables_html += single_table
        if not is_last_chunk: all_tables_html += "<div style='page-break-after: always; margin-bottom: 30px;'></div>"

    table_full_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Print</title><style>@page {{ size: A4 landscape; margin: 4mm; }} body {{ font-family: 'Sarabun', Tahoma, sans-serif; margin: 0; padding: 5px; background-color: #FFFFFF; color: #000000; -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }} .header-title {{ text-align: center; font-size: 16px; font-weight: bold; margin-bottom: 4px; }} .banner {{ text-align: center; font-size: 14px; font-weight: bold; padding: 5px; margin-bottom: 4px; }} table {{ width: 100%; border-collapse: collapse; }} th, td {{ font-family: 'Sarabun', Tahoma, sans-serif; color: #000000 !important; }}</style></head><body>{all_tables_html}</body></html>"
    return raw_data, excel_buffer.getvalue(), json.dumps(table_full_html), engine_lbl_list, engine_details_t2
@st.cache_data(show_spinner=False)
def process_t3(b_id_t3, start_date_str, end_date_str, allowed_tuple, selected_branch_display):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            if b_id_t3 == "ทั้งหมด":
                sql = "SELECT r.*, b.branch_name FROM boiler_pressure_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql, (start_date_str, end_date_str))
            elif b_id_t3 == "รวมเฉพาะที่มีสิทธิ์":
                placeholders = ', '.join(['%s'] * len(allowed_tuple))
                sql = f"SELECT r.*, b.branch_name FROM boiler_pressure_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.branch_id IN ({placeholders}) AND r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql, allowed_tuple + (start_date_str, end_date_str))
            else:
                sql = "SELECT r.*, b.branch_name FROM boiler_pressure_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.branch_id = %s AND r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql, (b_id_t3, start_date_str, end_date_str))
            raw_data = cur.fetchall()
        conn.close()
    except Exception as e:
        st.error(f"⚠️ เกิดข้อผิดพลาดในการดึงข้อมูล: {e}")
        return [], None, None

    if not raw_data: return [], None, None
    seen_b_keys, branches_config = set(), []
    for r in raw_data:
        b_name = str(r.get('branch_name') or '').strip()
        try: b_code = b_name.split("สาขา ")[1].split()[0]
        except: b_code = b_name or "WU"
        if b_code.lower() not in seen_b_keys:
            seen_b_keys.add(b_code.lower())
            branches_config.append({"code": b_code, "full_name": b_name})
    if not branches_config: branches_config = [{"code": "WU", "full_name": selected_branch_display}]

    start_dt = pd.to_datetime(start_date_str)
    days_in_month = calendar.monthrange(start_dt.year, start_dt.month)[1]
    matrix_data = {d: {b['code']: {'total': 0.0, 'drop': 0.0, 'pm': 0.0, 'non_pm': 0.0, 'remark': '', 'has_data': False} for b in branches_config} for d in range(1, days_in_month + 1)}
    for r in raw_data:
        d_num = pd.to_datetime(r['record_date']).day
        b_name = str(r.get('branch_name') or '').strip()
        try: b_code = b_name.split("สาขา ")[1].split()[0]
        except: b_code = b_name or "WU"
        for b in branches_config:
            if b['code'].lower() == b_code.lower():
                matrix_data[d_num][b['code']] = {'total': float(r.get('total_count') or 0.0), 'drop': float(r.get('total_drop') or 0.0), 'pm': float(r.get('pm_drop') or 0.0), 'non_pm': float(r.get('non_pm_drop') or 0.0), 'remark': str(r.get('remark') or '').strip(), 'has_data': True}
                break
    branch_totals = {}
    for b in branches_config:
        b_code = b['code']
        tot_c = sum([matrix_data[d][b_code]['total'] for d in range(1, days_in_month + 1)])
        tot_d = sum([matrix_data[d][b_code]['drop'] for d in range(1, days_in_month + 1)])
        pm_d = sum([matrix_data[d][b_code]['pm'] for d in range(1, days_in_month + 1)])
        npm_d = sum([matrix_data[d][b_code]['non_pm'] for d in range(1, days_in_month + 1)])
        branch_totals[b_code] = {'total': tot_c, 'drop': tot_d, 'pm': pm_d, 'non_pm': npm_d, 'p_tot': f"{(tot_d/tot_c*100):.2f}%" if tot_c>0 else "0.00%", 'p_pm': f"{(pm_d/tot_c*100):.2f}%" if tot_c>0 else "0.00%", 'p_npm': f"{(npm_d/tot_c*100):.2f}%" if tot_c>0 else "0.00%"}

    start_dt = pd.to_datetime(start_date_str)
    month_str = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"][start_dt.month - 1]
    year_buddhist = start_dt.year + 543

    wb = Workbook()
    ws = wb.active
    ws.title = "Pressure Report Matrix"
    ws.views.sheetView[0].showGridLines = True
    font_title, font_banner, font_header, font_body, font_body_bold, font_total, font_total_red, font_total_blue = Font(name="Sarabun", size=14, bold=True, color="FFFFFF"), Font(name="Sarabun", size=12, bold=True, color="334155"), Font(name="Sarabun", size=10, bold=True, color="1E293B"), Font(name="Sarabun", size=9, color="1E293B"), Font(name="Sarabun", size=9, bold=True, color="1E293B"), Font(name="Sarabun", size=10, bold=True, color="1E293B"), Font(name="Sarabun", size=10, bold=True, color="D32F2F"), Font(name="Sarabun", size=10, bold=True, color="1976D2")
    align_center, align_right, align_left = Alignment(horizontal="center", vertical="center", wrap_text=True), Alignment(horizontal="right", vertical="center"), Alignment(horizontal="left", vertical="center")
    bd_color = "94A3B8"
    thin_border = Border(left=Side(style="thin", color=bd_color), right=Side(style="thin", color=bd_color), top=Side(style="thin", color=bd_color), bottom=Side(style="thin", color=bd_color))
    fill_title, fill_banner_bg, fill_header_1, fill_header_2, fill_drop_bg, fill_total_bg = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid"), PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"), PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"), PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"), PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid"), PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    last_col_letter = get_column_letter(1 + (len(branches_config) * 8))
    ws.merge_cells(f"A1:{last_col_letter}1"); ws["A1"] = "รายงานการตกของ แรงดันไอน้ำปลายทาง ของบอยเลอร์ แบบเปรียบเทียบ"; ws["A1"].font, ws["A1"].alignment, ws["A1"].fill, ws["A1"].border = font_title, align_center, fill_title, thin_border
    ws.merge_cells(f"A2:{last_col_letter}2"); ws["A2"] = f"ประจำเดือน {month_str} {year_buddhist}"; ws["A2"].font, ws["A2"].alignment, ws["A2"].fill, ws["A2"].border = font_banner, align_center, fill_banner_bg, thin_border
    ws.merge_cells("A3:A5"); ws["A3"] = "วันที่"; ws["A3"].font, ws["A3"].alignment, ws["A3"].fill, ws["A3"].border = font_header, align_center, fill_header_1, thin_border
    ws["A4"].border, ws["A5"].border = thin_border, thin_border

    col_idx = 2
    for b in branches_config:
        st_c, en_c = get_column_letter(col_idx), get_column_letter(col_idx + 7)
        ws.merge_cells(f"{st_c}3:{en_c}3"); ws[f"{st_c}3"] = b["code"]; ws[f"{st_c}3"].font, ws[f"{st_c}3"].alignment, ws[f"{st_c}3"].fill, ws[f"{st_c}3"].border = font_header, align_center, fill_header_1, thin_border
        for c in range(col_idx, col_idx + 8): ws[f"{get_column_letter(c)}3"].border = thin_border
        c_cnt, c_drp_st, c_drp_en, c_pct_st, c_pct_en, c_rem = get_column_letter(col_idx), get_column_letter(col_idx + 1), get_column_letter(col_idx + 3), get_column_letter(col_idx + 4), get_column_letter(col_idx + 6), get_column_letter(col_idx + 7)
        ws[f"{c_cnt}4"] = "นับทั้งหมด"; ws.merge_cells(f"{c_drp_st}4:{c_drp_en}4"); ws[f"{c_drp_st}4"] = "แรงดันตก (ครั้ง)"; ws.merge_cells(f"{c_pct_st}4:{c_pct_en}4"); ws[f"{c_pct_st}4"] = "แรงดันตก (%)"; ws[f"{c_rem}4"] = "หมายเหตุ"
        for c_i in range(col_idx, col_idx + 8): cell_h = ws[f"{get_column_letter(c_i)}4"]; cell_h.font, cell_h.alignment, cell_h.fill, cell_h.border = font_header, align_center, fill_header_2, thin_border
        sub_cols = ["จำนวน\n(ครั้ง)", "ตกทั้งหมด", "ตกตาม\nเงื่อนไขPM", "เกินนอกเหนือ\nจากการPM", "ตกทั้งหมด", "ตกตาม\nเงื่อนไขPM", "เกินนอกเหนือ\nจากการPM", ""]
        for i, sh in enumerate(sub_cols): cell_ref = f"{get_column_letter(col_idx + i)}5"; ws[cell_ref] = sh; ws[cell_ref].font, ws[cell_ref].alignment, ws[cell_ref].fill, ws[cell_ref].border = Font(name="Sarabun", size=9, bold=True, color="475569"), align_center, fill_banner_bg, thin_border
        col_idx += 8

    current_row = 6
    for day in range(1, days_in_month + 1):
        ws[f"A{current_row}"] = day; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].border = font_body_bold, align_center, thin_border
        c_idx = 2
        for b in branches_config:
            item = matrix_data[day][b['code']]
            if item['has_data']:
                ws[f"{get_column_letter(c_idx)}{current_row}"].value = item['total']; ws[f"{get_column_letter(c_idx+1)}{current_row}"].value = item['drop']; ws[f"{get_column_letter(c_idx+2)}{current_row}"].value = item['pm']; ws[f"{get_column_letter(c_idx+3)}{current_row}"].value = item['non_pm']; ws[f"{get_column_letter(c_idx+4)}{current_row}"].value = f"{(item['drop']/item['total']*100):.2f}%" if item['total']>0 else "0.00%"; ws[f"{get_column_letter(c_idx+5)}{current_row}"].value = f"{(item['pm']/item['total']*100):.2f}%" if item['total']>0 else "0.00%"; ws[f"{get_column_letter(c_idx+6)}{current_row}"].value = f"{(item['non_pm']/item['total']*100):.2f}%" if item['total']>0 else "0.00%"; ws[f"{get_column_letter(c_idx+7)}{current_row}"].value = item['remark']
            for offset in range(8):
                cell = ws[f"{get_column_letter(c_idx+offset)}{current_row}"]
                cell.font, cell.border = font_body, thin_border
                if offset < 4 and cell.value != "": cell.alignment, cell.fill, cell.font = (align_center if offset==0 else align_right), fill_drop_bg, font_body_bold
                elif offset >= 4 and offset < 7 and cell.value != "": cell.alignment = align_right
                elif offset == 7: cell.alignment = align_left
            c_idx += 8
        current_row += 1

    ws[f"A{current_row}"] = "รวม"; ws[f"A{current_row}"].font, ws[f"A{current_row}"].alignment, ws[f"A{current_row}"].fill, ws[f"A{current_row}"].border = font_total, align_center, fill_total_bg, thin_border
    c_idx = 2
    for b in branches_config:
        bt = branch_totals[b['code']]
        ws[f"{get_column_letter(c_idx)}{current_row}"] = bt['total'] if bt['total'] > 0 else "-"
        ws[f"{get_column_letter(c_idx+1)}{current_row}"] = bt['drop'] if bt['drop'] > 0 else "-"
        ws[f"{get_column_letter(c_idx+2)}{current_row}"] = bt['pm'] if bt['pm'] > 0 else "-"
        ws[f"{get_column_letter(c_idx+3)}{current_row}"] = bt['non_pm'] if bt['non_pm'] > 0 else "-"
        ws[f"{get_column_letter(c_idx+4)}{current_row}"] = bt['p_tot']; ws[f"{get_column_letter(c_idx+5)}{current_row}"] = bt['p_pm']; ws[f"{get_column_letter(c_idx+6)}{current_row}"] = bt['p_npm']
        for offset in range(8): 
            cell = ws[f"{get_column_letter(c_idx+offset)}{current_row}"]; cell.fill, cell.border = fill_total_bg, thin_border
        ws[f"{get_column_letter(c_idx)}{current_row}"].font, ws[f"{get_column_letter(c_idx)}{current_row}"].alignment = font_total, align_center
        for offset in range(1, 4): ws[f"{get_column_letter(c_idx+offset)}{current_row}"].font, ws[f"{get_column_letter(c_idx+offset)}{current_row}"].alignment = font_total, align_right
        ws[f"{get_column_letter(c_idx+4)}{current_row}"].font, ws[f"{get_column_letter(c_idx+4)}{current_row}"].alignment = font_total_red, align_right
        ws[f"{get_column_letter(c_idx+5)}{current_row}"].font, ws[f"{get_column_letter(c_idx+5)}{current_row}"].alignment = font_total, align_right
        ws[f"{get_column_letter(c_idx+6)}{current_row}"].font, ws[f"{get_column_letter(c_idx+6)}{current_row}"].alignment = font_total_blue, align_right
        c_idx += 8

    ws.column_dimensions['A'].width = 8
    for c in range(2, col_idx): ws.column_dimensions[get_column_letter(c)].width = 13
    excel_buffer = io.BytesIO(); wb.save(excel_buffer); excel_data = excel_buffer.getvalue()

    header_row1 = "".join([f'<th colspan="8" style="background-color:#E2E8F0;color:#1E293B;border:1px solid #CBD5E1;padding:6px;font-size:12px;text-align:center;">{b["code"]}</th>' for b in branches_config])
    header_row2 = "".join(['<th style="background-color:#F1F5F9;color:#1E293B;border:1px solid #CBD5E1;padding:4px;font-size:10px;">นับทั้งหมด</th><th colspan="3" style="background-color:#F1F5F9;color:#1E293B;border:1px solid #CBD5E1;padding:4px;font-size:10px;">แรงดันตก (ครั้ง)</th><th colspan="3" style="background-color:#F1F5F9;color:#1E293B;border:1px solid #CBD5E1;padding:4px;font-size:10px;">แรงดันตก (%)</th><th style="background-color:#F1F5F9;color:#1E293B;border:1px solid #CBD5E1;padding:4px;font-size:10px;">หมายเหตุ</th>' for _ in branches_config])
    header_row3 = "".join(['<th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:35px;">จำนวน<br>(ครั้ง)</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:35px;">ตกทั้งหมด</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:45px;">ตกตาม<br>เงื่อนไขPM</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:45px;">เกินนอกเหนือ<br>จากการPM</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:35px;">ตกทั้งหมด</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:45px;">ตกตาม<br>เงื่อนไขPM</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:45px;">เกินนอกเหนือ<br>จากการPM</th><th style="background-color:#F8FAFC;color:#475569;border:1px solid #CBD5E1;padding:4px;font-size:9px;width:100px;"></th>' for _ in branches_config])

    body_rows = ""
    for day in range(1, days_in_month + 1):
        body_rows += f'<tr><td style="background-color:#F8FAFC; border:1px solid #CBD5E1;padding:4px;text-align:center;font-size:10px;font-weight:bold;color:#1E293B;">{day}</td>'
        for b in branches_config:
            item = matrix_data[day][b['code']]
            if item['has_data']:
                c_tot_s, c_drp_s = f"{item['total']:,.2f}", f"{item['drop']:,.2f}"
                c_pm_s, c_npm_s = f"{int(item['pm'])}", f"{int(item['non_pm'])}"
                p_drp_s = f"{(item['drop'] / item['total'] * 100):.2f}%" if item['total'] > 0 else "0.00%"
                p_pm_s = f"{(item['pm'] / item['total'] * 100):.2f}%" if item['total'] > 0 else "0.00%"
                p_npm_s = f"{(item['non_pm'] / item['total'] * 100):.2f}%" if item['total'] > 0 else "0.00%"
                rem_s = item['remark']
            else: c_tot_s, c_drp_s, c_pm_s, c_npm_s, p_drp_s, p_pm_s, p_npm_s, rem_s = "", "", "", "", "", "", "", ""
            body_rows += f'<td style="background-color:#FFFBEB;border:1px solid #CBD5E1;padding:4px;text-align:center;font-size:10px;font-weight:bold;">{c_tot_s}</td><td style="background-color:#FFFBEB;border:1px solid #CBD5E1;padding:4px;text-align:right;font-size:10px;font-weight:bold;">{c_drp_s}</td><td style="background-color:#FFFBEB;border:1px solid #CBD5E1;padding:4px;text-align:right;font-size:10px;font-weight:bold;">{c_pm_s}</td><td style="background-color:#FFFBEB;border:1px solid #CBD5E1;padding:4px;text-align:right;font-size:10px;font-weight:bold;">{c_npm_s}</td><td style="border:1px solid #CBD5E1;padding:4px;text-align:right;font-size:10px;">{p_drp_s}</td><td style="border:1px solid #CBD5E1;padding:4px;text-align:right;font-size:10px;">{p_pm_s}</td><td style="border:1px solid #CBD5E1;padding:4px;text-align:right;font-size:10px;">{p_npm_s}</td><td style="border:1px solid #CBD5E1;padding:4px;text-align:left;font-size:10px;color:#475569;">{rem_s}</td>'
        body_rows += '</tr>'

    total_row_html = '<tr><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:center;font-size:11px;font-weight:bold;color:#1E293B;">รวม</td>'
    for b in branches_config:
        bt = branch_totals[b['code']]
        t_cnt_s = f"{int(bt['total']):,}" if bt['total'] > 0 else "-"
        t_drp_s = f"{int(bt['drop']):,}" if bt['drop'] > 0 else "-"
        t_pm_s = f"{int(bt['pm']):,}" if bt['pm'] > 0 else "-"
        t_npm_s = f"{int(bt['non_pm']):,}" if bt['non_pm'] > 0 else "-"
        total_row_html += f'<td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:center;font-size:10px;font-weight:bold;">{t_cnt_s}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:right;font-size:10px;font-weight:bold;">{t_drp_s}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:right;font-size:10px;font-weight:bold;">{t_pm_s}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:right;font-size:10px;font-weight:bold;">{t_npm_s}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:right;font-size:10px;font-weight:bold;color:#D32F2F;">{bt["p_tot"]}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:right;font-size:10px;font-weight:bold;color:#1E293B;">{bt["p_pm"]}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:right;font-size:10px;font-weight:bold;color:#1976D2;">{bt["p_npm"]}</td><td style="background-color:#F1F5F9;border:1px solid #CBD5E1;padding:6px;text-align:left;font-size:10px;"></td>'
    total_row_html += '</tr>'

    table_full_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Print</title><style>@page {{ size: A4 portrait; margin: 5mm 8mm; }} body {{ font-family: 'Sarabun', Tahoma, sans-serif; margin: 0; padding: 0; background-color: #FFFFFF; color: #000000; -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }} .header-banner1 {{ text-align: center; font-size: 14px; font-weight: bold; margin-bottom: 2px; padding: 4px; background-color: #2D2D2D !important; color: #FFFFFF !important; }} .header-banner2 {{ text-align: center; font-size: 12px; font-weight: bold; padding: 3px; margin-bottom: 4px; background-color: #F3F4F6 !important; }} table {{ width: 100%; border-collapse: collapse; }} th, td {{ font-family: 'Sarabun', Tahoma, sans-serif; color: #000000 !important; border: 1px solid #CBD5E1 !important; padding: 1.5px 2px !important; font-size: 8.5px !important; line-height: 1.05 !important; text-align: center; }}</style></head><body><div class='header-banner1'>รายงานการตกของ แรงดันไอน้ำปลายทาง ของบอยเลอร์ แบบเปรียบเทียบ</div><div class='header-banner2'>ประจำเดือน {month_str} {year_buddhist}</div><table><thead><tr><th rowspan='3' style='background-color:#F1F5F9;color:#1E293B;padding:4px;font-size:10px;width:35px;'>วันที่</th>{header_row1}</tr><tr>{header_row2}</tr><tr>{header_row3}</tr></thead><tbody>{body_rows}{total_row_html}</tbody></table></body></html>"
    return raw_data, excel_data, json.dumps(table_full_html)

@st.cache_data(show_spinner=False)
def process_t4(b_id_t4, start_date_str, end_date_str, allowed_tuple, selected_branch_display):
    try:
        conn = get_db_connection()
        with conn.cursor() as cur:
            std = {'p_cubft': 12.00, 'stm_mul': 0.1875, 'p_ton': 240.00, 'avg_p': 3.00, 'kg_ton': 300.00, 'kg_cubft': 15.00}
            cur.execute("SELECT * FROM std_settings LIMIT 1")
            s_res = cur.fetchone()
            if s_res:
                std['p_cubft'] = float(s_res.get('std_fuel_price_per_cubft', 12.00))
                std['stm_mul'] = float(s_res.get('std_steam_multiplier', 0.1875))
                std['p_ton'] = float(s_res.get('std_fuel_price_per_steam_ton', 240.00))
                std['avg_p'] = float(s_res.get('std_avg_pressure', 3.00))
                std['kg_ton'] = float(s_res.get('std_fuel_kg_per_steam_ton', 300.00))
                std['kg_cubft'] = float(s_res.get('std_fuel_kg_per_cubft', 15.00))

            if b_id_t4 == "ทั้งหมด":
                sql = "SELECT r.*, b.branch_name FROM boiler_fuel_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql, (start_date_str, end_date_str))
                raw_data = cur.fetchall()
                sql_oven = "SELECT o.*, b.branch_name FROM daily_wood_oven_records o LEFT JOIN branches b ON o.branch_id = b.id WHERE o.record_date BETWEEN %s AND %s ORDER BY o.record_date ASC"
                cur.execute(sql_oven, (start_date_str, end_date_str))
                oven_raw = cur.fetchall()
            elif b_id_t4 == "รวมเฉพาะที่มีสิทธิ์":
                placeholders = ', '.join(['%s'] * len(allowed_tuple))
                sql = f"SELECT r.*, b.branch_name FROM boiler_fuel_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.branch_id IN ({placeholders}) AND r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql, allowed_tuple + (start_date_str, end_date_str))
                raw_data = cur.fetchall()
                sql_oven = f"SELECT o.*, b.branch_name FROM daily_wood_oven_records o LEFT JOIN branches b ON o.branch_id = b.id WHERE o.branch_id IN ({placeholders}) AND o.record_date BETWEEN %s AND %s ORDER BY o.record_date ASC"
                cur.execute(sql_oven, allowed_tuple + (start_date_str, end_date_str))
                oven_raw = cur.fetchall()
            else:
                sql = "SELECT r.*, b.branch_name FROM boiler_fuel_records r LEFT JOIN branches b ON r.branch_id = b.id WHERE r.branch_id = %s AND r.record_date BETWEEN %s AND %s ORDER BY r.record_date ASC"
                cur.execute(sql, (b_id_t4, start_date_str, end_date_str))
                raw_data = cur.fetchall()
                sql_oven = "SELECT o.*, b.branch_name FROM daily_wood_oven_records o LEFT JOIN branches b ON o.branch_id = b.id WHERE o.branch_id = %s AND o.record_date BETWEEN %s AND %s ORDER BY o.record_date ASC"
                cur.execute(sql_oven, (b_id_t4, start_date_str, end_date_str))
                oven_raw = cur.fetchall()
        conn.close()
    except Exception as e:
        st.error(f"⚠️ เกิดข้อผิดพลาดในการดึงข้อมูล: {e}")
        return [], [], None, None
    
    if not raw_data and not oven_raw: return [], [], None, None

    complete_set = set()
    if oven_raw and raw_data:
        oven_set = set((str(r.get('record_date')), str(r.get('branch_id'))) for r in oven_raw)
        b1_set = set((str(r.get('record_date')), str(r.get('branch_id'))) for r in raw_data if str(r.get('boiler_name')) == 'บอยเลอร์ 1')
        b2_set = set((str(r.get('record_date')), str(r.get('branch_id'))) for r in raw_data if str(r.get('boiler_name')) == 'บอยเลอร์ 2')
        complete_set = oven_set.intersection(b1_set).intersection(b2_set)
        oven_raw_filtered = [r for r in oven_raw if (str(r.get('record_date')), str(r.get('branch_id'))) in complete_set]
    else: oven_raw_filtered = []

    start_dt = pd.to_datetime(start_date_str)
    month_str = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"][start_dt.month - 1]
    year_buddhist = start_dt.year + 543
    master_summary, grouped_data = {}, {}

    for r in raw_data:
        b_name = str(r.get('boiler_name') or 'บอยเลอร์ 1')
        if b_name not in grouped_data: grouped_data[b_name] = []
        grouped_data[b_name].append(r)
        
        if (str(r.get('record_date')), str(r.get('branch_id'))) in complete_set:
            bn, dt = str(r.get('branch_name') or '-'), r['record_date']
            k = (bn, dt)
            if k not in master_summary: master_summary[k] = {'oven':0, 'wood':0, 'press':0, 'press_cnt':0, 'sw':0, 'ww':0, 'www':0, 'sp':0, 'wp':0, 'wwp':0, 'stm':0, 'hrs':0}
            master_summary[k]['sw'] += float(r.get('sawdust_weight') or 0); master_summary[k]['ww'] += float(r.get('wood_weight') or 0); master_summary[k]['www'] += float(r.get('waste_wood_weight') or 0)
            master_summary[k]['sp'] += float(r.get('sawdust_price') or 0); master_summary[k]['wp'] += float(r.get('wood_price') or 0); master_summary[k]['wwp'] += float(r.get('waste_wood_price') or 0)
            master_summary[k]['stm'] += float(r.get('steam_production') or 0); master_summary[k]['hrs'] += float(r.get('working_hours') or 0)

    for o in oven_raw_filtered:
        bn, dt = str(o.get('branch_name') or '-'), o['record_date']
        k = (bn, dt)
        if k not in master_summary: master_summary[k] = {'oven':0, 'wood':0, 'press':0, 'press_cnt':0, 'sw':0, 'ww':0, 'www':0, 'sp':0, 'wp':0, 'wwp':0, 'stm':0, 'hrs':0}
        master_summary[k]['oven'] = float(o.get('oven_qty') or 0); master_summary[k]['wood'] = float(o.get('wood_out_cubft') or 0); master_summary[k]['press'] += float(o.get('avg_terminal_pressure') or 0); master_summary[k]['press_cnt'] += 1

    wb = Workbook()
    wb.remove(wb.active)
    f_b, f_w, f_n, f_red, f_bold_ul = Font(name="Tahoma", size=10, bold=True), Font(name="Tahoma", size=10, bold=True, color="FFFFFF"), Font(name="Tahoma", size=10), Font(name="Tahoma", size=10, color="FF0000"), Font(name="Tahoma", size=10, bold=True, underline="single")
    al_c, al_r = Alignment(horizontal="center", vertical="center", wrap_text=True), Alignment(horizontal="right", vertical="center")
    tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    fill_blue, fill_yellow, fill_green, fill_peach = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid"), PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"), PatternFill(start_color="00B050", end_color="00B050", fill_type="solid"), PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    html_tables, boilers_html_dict = "", {}

    ws_m = wb.create_sheet(title=f"สรุปรวม 1+2")
    ws_m.merge_cells("A1:AD1"); ws_m["A1"] = f"รายงานการใช้เชื้อเพลิง {selected_branch_display} บอยเลอร์ 1+2 เดือน {month_str} {year_buddhist}"; ws_m["A1"].font, ws_m["A1"].fill, ws_m["A1"].alignment = f_w, fill_blue, al_c
    headers_m = [("A2:A3","วันที่"), ("B2:B3","จำนวน\nเตาที่อบ"), ("C2:C3","ไม้อบออก\n(ลบ.ฟ.)"), ("D2:F2","เชื้อเพลิงใช้(ตัน/วัน)"), ("G2:G3","รวมน.น.\nเชื้อเพลิง"), ("H2:K2","ค่าเชื้อเพลิง(บาท)"), ("L2:N2","ค่าเชื้อเพลิง(บาท/ลบ.ฟ.)"), ("O2:R2","ผลิตไอน้ำ"), ("S2:U2","ค่าเชื้อเพลิง(บาท/ตันไอน้ำ)"), ("V2:X2","แรงดันปลายทางเฉลี่ย"), ("Y2:AA2","เชื้อเพลิงใช้(กก./ตันไอน้ำ)"), ("AB2:AD2","เชื้อเพลิงใช้(กก./ลบ.ฟ.)")]
    for cell_r, val in headers_m: ws_m.merge_cells(cell_r); ws_m[cell_r.split(":")[0]] = val
    h3_m = ["", "", "", "ขี้เลื่อย", "ปีกไม้", "เศษไม้เสีย", "", "ขี้เลื่อย", "ปีกไม้", "เศษไม้เสีย", "รวมเป็นเงิน", "STD", "ผลงาน", "ผลต่าง", "ตัน/วัน", "STD\nตัน/ชม.", "ผลงาน", "ผลต่าง", "STD", "ผลงาน", "ผลต่าง", "STD", "ผลงาน", "ผลต่าง", "STD", "ผลงาน", "ผลต่าง", "STD", "ผลงาน", "ผลต่าง"]
    for idx, val in enumerate(h3_m, 1):
        if val: ws_m.cell(row=3, column=idx, value=val)
    for r_idx in [2, 3]:
        for c_idx in range(1, 31): c = ws_m.cell(row=r_idx, column=c_idx); c.font, c.alignment, c.border = f_b, al_c, tb

    html_headers_m = """<tr><th rowspan="2">วันที่</th><th rowspan="2">จำนวน<br>เตาที่อบ</th><th rowspan="2">ไม้อบออก<br>(ลบ.ฟ.)</th><th colspan="3">เชื้อเพลิงใช้(ตัน/วัน)</th><th rowspan="2">รวมน.น.<br>เชื้อเพลิง</th><th colspan="4">ค่าเชื้อเพลิง(บาท)</th><th colspan="3">ค่าเชื้อเพลิง(บาท/ลบ.ฟ.)</th><th colspan="4">ผลิตไอน้ำ</th><th colspan="3">ค่าเชื้อเพลิง(บาท/ตันไอน้ำ)</th><th colspan="3">แรงดันปลายทางเฉลี่ย</th><th colspan="3">เชื้อเพลิงใช้(กก./ตันไอน้ำ)</th><th colspan="3">เชื้อเพลิงใช้(กก./ลบ.ฟ.)</th></tr><tr><th>ขี้เลื่อย</th><th>ปีกไม้</th><th>เศษไม้เสีย</th><th>ขี้เลื่อย</th><th>ปีกไม้</th><th>เศษไม้เสีย</th><th>รวมเป็นเงิน</th><th>STD</th><th>ผลงาน</th><th>ผลต่าง</th><th>ตัน/วัน</th><th>STD<br>ตัน/ชม.</th><th>ผลงาน</th><th>ผลต่าง</th><th>STD</th><th>ผลงาน</th><th>ผลต่าง</th><th>STD</th><th>ผลงาน</th><th>ผลต่าง</th><th>STD</th><th>ผลงาน</th><th>ผลต่าง</th><th>STD</th><th>ผลงาน</th><th>ผลต่าง</th></tr>"""

    cur_row = 4; body_html_m = ""
    s_ov = s_wo = s_sw = s_ww = s_www = s_tw = s_sp = s_wp = s_wwp = s_tp = s_stm = s_hrs = s_prs = s_pcnt = 0
    sorted_master_data = sorted(master_summary.items(), key=lambda x: (x[0][1], x[0][0]))

    for k, v in sorted_master_data:
        branch_full, dt_val = k[0], k[1]
        try: b_short = branch_full.split("สาขา ")[1].split()[0]
        except: b_short = branch_full
        if not b_short or b_short == '-': b_short = "N/A"
        dt = pd.to_datetime(dt_val)
        date_str = f"{dt.day:02d}/{dt.month:02d}/{(dt.year+543)%100}({b_short})"
        
        tw, tp = v['sw'] + v['ww'] + v['www'], v['sp'] + v['wp'] + v['wwp']
        avg_p = v['press'] / v['press_cnt'] if v['press_cnt'] > 0 else 0
        act_p_cub = tp / v['wood'] if v['wood'] > 0 else 0
        std_stm_hr = v['oven'] * std['stm_mul']
        act_stm_hr = v['stm'] / v['hrs'] if v['hrs'] > 0 else 0
        act_p_ton = tp / v['stm'] if v['stm'] > 0 else 0
        act_kg_ton = (tw * 1000) / v['stm'] if v['stm'] > 0 else 0
        act_kg_cub = (tw * 1000) / v['wood'] if v['wood'] > 0 else 0
        
        vals = [date_str, v['oven'], v['wood'], v['sw'], v['ww'], v['www'], tw, v['sp'], v['wp'], v['wwp'], tp, std['p_cubft'], act_p_cub, std['p_cubft'] - act_p_cub, v['stm'], std_stm_hr, act_stm_hr, act_stm_hr - std_stm_hr, std['p_ton'], act_p_ton, std['p_ton'] - act_p_ton, std['avg_p'], avg_p, avg_p - std['avg_p'], std['kg_ton'], act_kg_ton, std['kg_ton'] - act_kg_ton, std['kg_cubft'], act_kg_cub, std['kg_cubft'] - act_kg_cub]
        s_ov+=v['oven']; s_wo+=v['wood']; s_sw+=v['sw']; s_ww+=v['ww']; s_www+=v['www']; s_tw+=tw; s_sp+=v['sp']; s_wp+=v['wp']; s_wwp+=v['wwp']; s_tp+=tp; s_stm+=v['stm']; s_hrs+=v['hrs']; s_prs+=v['press']; s_pcnt+=v['press_cnt']

        body_html_m += "<tr>"
        for idx, val in enumerate(vals, 1):
            c = ws_m.cell(row=cur_row, column=idx, value=val); c.border, c.alignment = tb, al_c if idx==1 else al_r
            is_diff_col = idx in [14, 18, 21, 24, 27, 30]; is_std_col = idx in [12, 16, 19, 22, 25, 28]
            c.font = f_red if is_diff_col and val < 0 else f_n
            if is_std_col: c.fill, c.font = fill_green, f_w
            elif idx in [2, 3, 23]: c.fill = fill_yellow 
            if isinstance(val, (int, float)): c.number_format = '#,##0.00'
            bg = "background-color:#00B050;color:white;" if is_std_col else ("background-color:#FFFF00;" if idx in [2, 3, 23] else "")
            tc = "color:red;" if is_diff_col and val < 0 else ""
            fw = "font-weight:bold;" if idx in [7,11] else ""
            body_html_m += f"<td style='{bg}{tc}{fw}text-align:{'center' if idx==1 else 'right'};'>{val if idx==1 else f'{val:,.2f}'}</td>"
        body_html_m += "</tr>"
        cur_row += 1

    avg_p_cub = s_tp / s_wo if s_wo > 0 else 0
    sum_std_stm = s_ov * std['stm_mul']
    avg_stm_hr = s_stm / s_hrs if s_hrs > 0 else 0
    avg_p_ton = s_tp / s_stm if s_stm > 0 else 0
    avg_press = s_prs / s_pcnt if s_pcnt > 0 else 0
    avg_kg_ton = (s_tw * 1000) / s_stm if s_stm > 0 else 0
    avg_kg_cub = (s_tw * 1000) / s_wo if s_wo > 0 else 0

    f1_vals = ["รวมทั้งหมด", s_ov, s_wo, s_sw, s_ww, s_www, s_tw, s_sp, s_wp, s_wwp, s_tp, "", avg_p_cub, std['p_cubft'] - avg_p_cub, s_stm, sum_std_stm, avg_stm_hr, avg_stm_hr - sum_std_stm, "", avg_p_ton, std['p_ton'] - avg_p_ton, "", avg_press, avg_press - std['avg_p'], "", avg_kg_ton, std['kg_ton'] - avg_kg_ton, "", avg_kg_cub, std['kg_cubft'] - avg_kg_cub]
    f_html_m = "<tr style='font-weight:bold; background-color:#FCE4D6;'>"
    for idx, val in enumerate(f1_vals, 1):
        c = ws_m.cell(row=cur_row, column=idx, value=val); c.border, c.font, c.alignment, c.fill = tb, f_b, al_r if idx > 1 else al_c, fill_peach
        is_diff_col = idx in [14, 18, 21, 24, 27, 30]
        if is_diff_col and val != "" and val < 0: c.font = f_red
        if isinstance(val, (int, float)): c.number_format = '#,##0.00'
        tc = "color:red;" if is_diff_col and val != "" and val < 0 else ""
        f_html_m += f"<td style='{tc}text-align:{'center' if idx==1 else 'right'};'>{val if val == '' or idx==1 else f'{val:,.2f}'}</td>"
    f_html_m += "</tr>"
    for i in range(1, 31): ws_m.column_dimensions[get_column_letter(i)].width = 10
    ws_m.column_dimensions['A'].width = 14

    master_html = f"<table style='margin-bottom: 30px;'><thead><tr><th colspan='30' class='header-main'>รายงานการใช้เชื้อเพลิง {selected_branch_display} บอยเลอร์ 1+2 เดือน {month_str} {year_buddhist}</th></tr>{html_headers_m}</thead><tbody>{body_html_m}{f_html_m}</tbody></table>"
    boilers_html_dict[f"🌟 รวมสาขา {selected_branch_display} (บอยเลอร์ 1+2)"] = f"<!DOCTYPE html><html><head><meta charset='utf-8'><style>@page {{ size: A4 landscape; margin: 4mm; }} * {{ -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }} body {{ font-family: 'Sarabun', Tahoma, sans-serif; font-size: 8px; color: #000; margin: 0; padding: 2px; }} table {{ width: 100%; border-collapse: collapse; }} th, td {{ border: 1px solid #000; padding: 1.5px 2px !important; color: #000 !important; font-size: 7.5px !important; line-height: 1.05 !important; white-space: nowrap !important; }} th {{ background-color: #F8F9FA; text-align: center; font-weight: bold; white-space: normal !important; }} .header-main {{ background-color: #0070C0; color: white !important; font-size: 12px; padding: 4px; }}</style></head><body>{master_html}</body></html>"
    html_tables += master_html

    for boiler_name in sorted(grouped_data.keys()):
        data_list = sorted(grouped_data[boiler_name], key=lambda x: (str(x.get('record_date') or ''), str(x.get('branch_name') or '')))
        ws = wb.create_sheet(title=boiler_name)
        ws.merge_cells("A1:O1"); ws["A1"] = f"รายงานการใช้เชื้อเพลิง {selected_branch_display} {boiler_name} เดือน {month_str} {year_buddhist}"; ws["A1"].font, ws["A1"].fill, ws["A1"].alignment = f_w, fill_blue, al_c
        headers_merge = [("A2:A3","วันที่"), ("B2:D2","เชื้อเพลิงใช้(ตัน/วัน)"), ("E2:E3","รวม น.น.\nเชื้อเพลิง"), ("F2:I2","ค่าเชื้อเพลิง(บาท)"), ("J2:K2","ผลิตไอน้ำ"), ("L2:L3","ค่าเชื้อเพลิง\n(บาท/ตันไอน้ำ)"), ("M2:O2","เชื้อเพลิงใช้(กก./ตันไอน้ำ)")]
        for cell_r, val in headers_merge: ws.merge_cells(cell_r); ws[cell_r.split(":")[0]] = val
        h3 = ["", "ขี้เลื่อย", "ปีกไม้", "เศษไม้เสีย", "", "ขี้เลื่อย", "ปีกไม้", "เศษไม้เสีย", "รวมเป็นเงิน", "ตัน/วัน", "ตัน/ชม.", "", "STD", "ผลงาน", "ผลต่าง"]
        for idx, val in enumerate(h3, 1):
            if val: ws.cell(row=3, column=idx, value=val)
        for r_idx in [2, 3]:
            for c_idx in range(1, 16): c = ws.cell(row=r_idx, column=c_idx); c.font, c.alignment, c.border = f_b, al_c, tb
        html_headers = """<tr><th rowspan="2">วันที่</th><th colspan="3">เชื้อเพลิงใช้(ตัน/วัน)</th><th rowspan="2">รวม น.น.<br>เชื้อเพลิง</th><th colspan="4">ค่าเชื้อเพลิง(บาท)</th><th colspan="2">ผลิตไอน้ำ</th><th rowspan="2">ค่าเชื้อเพลิง<br>(บาท/ตันไอน้ำ)</th><th colspan="3">เชื้อเพลิงใช้(กก./ตันไอน้ำ)</th></tr><tr><th>ขี้เลื่อย</th><th>ปีกไม้</th><th>เศษไม้เสีย</th><th>ขี้เลื่อย</th><th>ปีกไม้</th><th>เศษไม้เสีย</th><th>รวมเป็นเงิน</th><th>ตัน/วัน</th><th>ตัน/ชม.</th><th>STD</th><th>ผลงาน</th><th>ผลต่าง</th></tr>"""

        cur_row = 4; body_html = ""
        sum_sw = sum_ww = sum_www = sum_tot_w = sum_sp = sum_wp = sum_wwp = sum_tot_p = sum_steam = sum_hrs = 0
        for r in data_list:
            dt = pd.to_datetime(r['record_date'])
            bn = str(r.get('branch_name') or '-')
            try: b_short = bn.split("สาขา ")[1].split()[0]
            except: b_short = bn
            if not b_short or b_short == '-': b_short = "N/A"
            date_str = f"{dt.day:02d}/{dt.month:02d}/{(dt.year+543)%100}({b_short})"
            
            s_w, w_w, ww_w = float(r.get('sawdust_weight') or 0), float(r.get('wood_weight') or 0), float(r.get('waste_wood_weight') or 0)
            s_p, w_p, ww_p = float(r.get('sawdust_price') or 0), float(r.get('wood_price') or 0), float(r.get('waste_wood_price') or 0)
            stm, hrs = float(r.get('steam_production') or 0), float(r.get('working_hours') or 0)
            t_w, t_p = (s_w + w_w + ww_w), (s_p + w_p + ww_p)
            stm_hr = stm / hrs if hrs > 0 else 0
            cost_ton = t_p / stm if stm > 0 else 0
            fuel_ton = (t_w * 1000) / stm if stm > 0 else 0
            diff = std['kg_ton'] - fuel_ton
            
            sum_sw+=s_w; sum_ww+=w_w; sum_www+=ww_w; sum_tot_w+=t_w; sum_sp+=s_p; sum_wp+=w_p; sum_wwp+=ww_p; sum_tot_p+=t_p; sum_steam+=stm; sum_hrs+=hrs
            row_vals = [date_str, s_w, w_w, ww_w, t_w, s_p, w_p, ww_p, t_p, stm, stm_hr, cost_ton, std['kg_ton'], fuel_ton, diff]
            
            body_html += "<tr>"
            for idx, val in enumerate(row_vals, 1):
                c = ws.cell(row=cur_row, column=idx, value=val); c.border, c.font, c.alignment = tb, (f_red if idx==15 and val<0 else f_n), (al_c if idx==1 else al_r)
                if idx in [2, 3, 4, 10]: c.fill = fill_yellow
                elif idx == 13: c.fill, c.font = fill_green, f_w
                if isinstance(val, (int, float)): c.number_format = '#,##0.00'
                bg = "background-color:#FFFF00;" if idx in [2,3,4,10] else ("background-color:#00B050;color:white;" if idx==13 else "")
                tc = "color:red;" if idx==15 and val<0 else ""
                fw = "font-weight:bold;" if idx in [5,9] else ""
                body_html += f"<td style='{bg}{tc}{fw}text-align:{'center' if idx==1 else 'right'};'>{val if idx==1 else f'{val:,.2f}'}</td>"
            body_html += "</tr>"
            cur_row += 1

        avg_stm_hr = sum_steam / sum_hrs if sum_hrs > 0 else 0
        avg_cost_ton = sum_tot_p / sum_steam if sum_steam > 0 else 0
        avg_fuel_ton = (sum_tot_w * 1000) / sum_steam if sum_steam > 0 else 0
        avg_diff = std['kg_ton'] - avg_fuel_ton
        
        ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=1)
        c_title = ws.cell(row=cur_row, column=1, value="รวมทั้งหมด"); c_title.alignment, c_title.fill, c_title.border, c_title.font = al_c, fill_peach, tb, f_bold_ul
        f1_vals = ["", sum_sw, sum_ww, sum_www, sum_tot_w, sum_sp, sum_wp, sum_wwp, sum_tot_p, sum_steam, avg_stm_hr, avg_cost_ton, std['kg_ton'], avg_fuel_ton, avg_diff]
        footer_html = f"<tr style='font-weight:bold; background-color:#FCE4D6;'><td style='text-align:center; text-decoration:underline;'>รวมทั้งหมด</td>"
        
        for idx in range(2, 16):
            val = f1_vals[idx-1]; c = ws.cell(row=cur_row, column=idx, value=val); c.border, c.font, c.alignment = tb, f_b, al_r
            if isinstance(val, (int, float)): c.number_format = '#,##0.00'
            if idx == 9: c.font = f_red
            if idx == 11: c.fill = fill_yellow
            if idx == 14: c.font = f_bold_ul
            if idx == 15 and val < 0: c.font = f_red
            bg = "background-color:#FFFF00;" if idx==11 else ""
            tc = "color:red;" if idx==9 or (idx==15 and val<0) else ""
            ul = "text-decoration:underline;" if idx==14 else ""
            footer_html += f"<td style='{bg}{tc}{ul}text-align:right;'>{f'{val:,.2f}' if isinstance(val, (int, float)) else ''}</td>"
        footer_html += "</tr>"
        
        for i in range(1, 16): ws.column_dimensions[get_column_letter(i)].width = 11
        ws.column_dimensions['A'].width = 14
        
        single_table_html = f"<table style='margin-bottom: 30px;'><thead><tr><th colspan='15' class='header-main'>รายงานการใช้เชื้อเพลิง {selected_branch_display} {boiler_name} เดือน {month_str} {year_buddhist}</th></tr>{html_headers}</thead><tbody>{body_html}{footer_html}</tbody></table>"
        html_tables += single_table_html
        boilers_html_dict[f"🔥 ข้อมูล {boiler_name}"] = f"<!DOCTYPE html><html><head><meta charset='utf-8'><style>@page {{ size: A4 landscape; margin: 8mm 5mm; }} * {{ -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }} body {{ font-family: 'Sarabun', Tahoma, sans-serif; color: #000; margin: 0; padding: 0; }} table {{ width: 100%; border-collapse: collapse; margin-top: 5px; }} th, td {{ border: 1px solid #000; padding: 2.5px 2px !important; color: #000 !important; font-size: 9px !important; line-height: 1.1 !important; white-space: nowrap !important; }} th {{ background-color: #F8F9FA; text-align: center; font-weight: bold; white-space: normal !important; }} .header-main {{ background-color: #0070C0; color: white !important; font-size: 14px; padding: 6px; }}</style></head><body>{single_table_html}</body></html>"

    full_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'><style>@page {{ size: A4 landscape; margin: 5mm; }} * {{ -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }} body {{ font-family: 'Sarabun', Tahoma, sans-serif; font-size:10px; color:#000; margin:0; padding:10px; }} table {{ width: 100%; border-collapse: collapse; }} th, td {{ border: 1px solid #000; padding: 4px; color:#000 !important; }} th {{ background-color: #F8F9FA; text-align: center; font-weight:bold; }} .header-main {{ background-color: #0070C0; color: white !important; font-size: 14px; padding: 6px; }}</style></head><body>{html_tables}</body></html>"
    final_json_data = { "full_print_html": full_html, "boilers": boilers_html_dict }
    excel_buffer = io.BytesIO(); wb.save(excel_buffer)
    return raw_data, oven_raw_filtered, excel_buffer.getvalue(), json.dumps(final_json_data)

# ==========================================================
# 📊 5. ฟังก์ชันหลักสำหรับ Render Report Tabs (รวม Dual View PC/Mobile)
# ==========================================================
def render_all_reports_module(user_branch_name, selected_sub_menu=None):
    
    st.markdown(""" 
        <style>
        header[data-testid="stHeader"] { display: none !important; }
        header { visibility: hidden !important; }
        #MainMenu { visibility: hidden !important; }
        .block-container { padding-top: 1rem !important; margin-top: -20px !important; }
        
        /* 📱 ระบบซ่อน/โชว์ อัจฉริยะ (Dual View) */
        @media (min-width: 769px) {
            div[data-testid="stVerticalBlock"]:has(.mobile-view-marker) { display: none !important; }
            div[data-testid="stVerticalBlock"]:has(.desktop-view-marker) { display: block !important; }
        }
        @media (max-width: 768px) {
            div[data-testid="stVerticalBlock"]:has(.desktop-view-marker) { display: none !important; }
            div[data-testid="stVerticalBlock"]:has(.mobile-view-marker) { display: block !important; }
        }

        /* เจาะจงลดขนาดปุ่มที่อยู่ในคอลัมน์ (เช่น ปุ่มในคอลัมน์ "จัดการ") */
        div[data-testid="column"] div.stButton > button {
            padding: 0px 5px !important;
            height: 35px !important;
            min-height: 35px !important;
            width: 100% !important;
            max-width: 45px !important; /* ล็อคความกว้างไม่ให้ยืดเกินไป */
            border-radius: 8px !important;
            margin: auto !important; /* จัดให้อยู่กึ่งกลาง */
        }
    
        /* ปรับขนาดไอคอนด้านในให้สมส่วน */
        div[data-testid="column"] div.stButton > button p {
            font-size: 16px !important;
            margin: 0 !important;
            line-height: 1 !important;
        }
        </style>
    """, unsafe_allow_html=True)

    raw_role = st.session_state.get('role_tab') or st.session_state.get('role') or 'user'
    current_role = str(raw_role).strip().lower()
    allowed_tabs_list = st.session_state.get('allowed_tabs', [])
    user_allowed_branches = st.session_state.get('allowed_branches', [str(st.session_state.branch_id)])

    key = "1"
    if selected_sub_menu:
        if "1." in selected_sub_menu: key = "1"
        elif "2." in selected_sub_menu: key = "2"
        elif "3." in selected_sub_menu: key = "3"
        elif "4." in selected_sub_menu: key = "4"

    if current_role not in ["admin", "reporter"] and key not in allowed_tabs_list:
        st.error("## ⏳ คุณไม่ได้รับสิทธิ์เข้าถึงหัวข้อนี้")
        return

    report_options = ["ทั้งหมดทุกสาขา"] + list(branch_dict.keys())
    current_tab_ctx = st.container()

    # =========================================================================
    # ⚙️ TAB 1: รายงานสรุปเครื่องจักร & เบรกดาวน์
    # =========================================================================
    if key == "1":
        with current_tab_ctx:
            st.subheader("📊 รายงานสรุปการทำงานและเบรกดาวน์เครื่องจักร")
            col_f1, col_f2, col_f3 = st.columns([1.5, 1.5, 2])
            with col_f1: start_date_t1 = st.date_input("ตั้งแต่วันที่", format="DD/MM/YYYY", value=pd.to_datetime("today").replace(day=1), key="s_d1")
            with col_f2: end_date_t1 = st.date_input("ถึงวันที่", format="DD/MM/YYYY", value=pd.to_datetime("today"), key="e_d1")
            with col_f3:
                if current_role in ["admin", "reporter"]:
                    b_label_t1 = st.selectbox("เลือกสาขา:", options=report_options, key="b_s1")
                    b_id_t1 = "ทั้งหมด" if b_label_t1 == "ทั้งหมดทุกสาขา" else branch_dict[b_label_t1]
                    selected_branch_display = b_label_t1 if b_label_t1 != "ทั้งหมดทุกสาขา" else "ทุกสาขา"
                elif len(user_allowed_branches) > 1:
                    allowed_options = ["ทั้งหมดที่มีสิทธิ์"] + [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
                    b_label_t1 = st.selectbox("เลือกสาขา:", options=allowed_options, key="b_s1")
                    b_id_t1 = "รวมเฉพาะที่มีสิทธิ์" if b_label_t1 == "ทั้งหมดที่มีสิทธิ์" else branch_dict[b_label_t1]
                    selected_branch_display = b_label_t1 if b_label_t1 != "ทั้งหมดที่มีสิทธิ์" else "ทุกสาขาที่มีสิทธิ์"
                else:
                    st.info(f"📍 สังกัด: {user_branch_name}")
                    b_id_t1 = st.session_state.branch_id
                    selected_branch_display = user_branch_name

            raw_data_t1, excel_data_t1, json_html_t1 = process_t1(b_id_t1, str(start_date_t1), str(end_date_t1), tuple(user_allowed_branches), selected_branch_display)

            if raw_data_t1:
                pk_col_t1 = list(raw_data_t1[0].keys())[0]
                paginated_t1 = render_custom_pagination(raw_data_t1, session_prefix="rep1")
                st.markdown("<br>", unsafe_allow_html=True)

                # =========================================================
                # 🖥️ DESKTOP VIEW: แสดงตารางแบบเดิม
                # =========================================================
                with st.container():
                    st.markdown('<div class="desktop-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                    # ปรับสัดส่วนคอลัมน์ใหม่ (ลบหมายเหตุออก เหลือ 8 คอลัมน์)
                    header_cols = st.columns([0.5, 1.1, 1.0, 1.0, 0.8, 1.1, 1.3, 1.2])
                    headers = ["No.", "วันที่", "สาขา", "ชื่อเครื่องจักร", "จำนวน", "ชม.ทำงาน", "ชม.เบรกดาวน์", "จัดการ"]
                    for col, header in zip(header_cols, headers): col.markdown(f"<span style='color:#64748B; font-weight:bold; font-size:14px;'>{header}</span>", unsafe_allow_html=True)
                    st.markdown("<hr style='margin: 0.2rem 0; border-color: #E2E8F0;'>", unsafe_allow_html=True)

                    with st.container(height=380, border=False):
                        start_num = (st.session_state.get("rep1_page", 1) - 1) * st.session_state.get("rep1_rows", 15)
                        for i, r in enumerate(paginated_t1):
                            seq_num = start_num + i + 1
                            rec_id = r[pk_col_t1]
                            cols = st.columns([0.5, 1.1, 1.0, 1.0, 0.8, 1.1, 1.3, 1.2])
                            cols[0].write(str(seq_num))
                            cols[1].write(pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-')
                            cols[2].markdown(f"<span class='branch-badge'>{r.get('branch_name') or '-'}</span>", unsafe_allow_html=True)
                            cols[3].write(str(r.get('machine_name') or '-'))
                            cols[4].write(str(r.get('machine_qty') or 0))
                            cols[5].write(f"{float(r.get('working_hours') or 0.0):.2f}")
                            cols[6].write(f"{float(r.get('breakdown_hours') or 0.0):.2f}")
                            # ลบ cols[7] หมายเหตุออก แล้วเลื่อนปุ่มจัดการมาแทน
                            
                            can_crud = current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)
                            if can_crud:
                                c_edit, c_del = cols[7].columns(2, gap="small")
                                if c_edit.button("✏️", key=f"e1_desk_{rec_id}", help="แก้ไขข้อมูล", use_container_width=True): update_record_dialog_t1(r, pk_col_t1)
                                if c_del.button("🗑️", key=f"d1_desk_{rec_id}", help="ลบรายการ", use_container_width=True): delete_record_dialog_t1(r, pk_col_t1)
                            else: cols[7].write("-")
                            st.markdown("<hr style='margin:0; border-color:#F1F5F9;'>", unsafe_allow_html=True)

                # =========================================================
                # 📱 MOBILE VIEW: แสดงตารางแบบใหม่ (Dataframe + แผงควบคุม)
                # =========================================================
                with st.container():
                    st.markdown('<div class="mobile-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                    if paginated_t1:
                        df_list = []
                        start_num = (st.session_state.get("rep1_page", 1) - 1) * st.session_state.get("rep1_rows", 15)
                        for i, r in enumerate(paginated_t1):
                            seq_num = start_num + i + 1
                            df_list.append({
                                "ลำดับ": seq_num,
                                "วันที่": pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-',
                                "สาขา": r.get('branch_name') or '-',
                                "ชื่อเครื่องจักร": r.get('machine_name') or '-',
                                "จำนวน": r.get('machine_qty') or 0,
                                "ชม.ทำงาน": f"{float(r.get('working_hours') or 0.0):.2f}",
                                "ชม.เบรกดาวน์": f"{float(r.get('breakdown_hours') or 0.0):.2f}",
                                "หมายเหตุ": r.get('remarks') or '-'
                            })
                        st.dataframe(pd.DataFrame(df_list), use_container_width=True, hide_index=True)

                        st.markdown("---")
                        st.markdown("<h5 style='color:#333333; font-weight:bold;'>🛠️ จัดการข้อมูล (แก้ไข / ลบรายการ)</h5>", unsafe_allow_html=True)
                        editable_records = [r for r in paginated_t1 if current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)]
                        
                        if editable_records:
                            col_sel, col_btn_e, col_btn_d = st.columns([2.5, 1, 1])
                            opts = {f"ID: {r[pk_col_t1]} | วันที่: {pd.to_datetime(r.get('record_date')).strftime('%d/%m/%Y')} | {r.get('machine_name')}": r for r in editable_records}
                            sel_lbl = col_sel.selectbox("เลือกลำดับข้อมูล", options=list(opts.keys()), label_visibility="collapsed", key="sel_crud_t1_mob")
                            sel_row = opts[sel_lbl]
                            
                            if col_btn_e.button("✏️", use_container_width=True, key=f"e1_btn_mob"): update_record_dialog_t1(sel_row, pk_col_t1)
                            if col_btn_d.button("🗑️", use_container_width=True, key=f"d1_btn_mob"): delete_record_dialog_t1(sel_row, pk_col_t1)
                        else:
                            st.info("ไม่มีสิทธิ์จัดการข้อมูล")

                # =========================================================
                # 🖨️ ปุ่มด้านล่าง (โชว์ทั้ง PC และ Mobile พร้อมแก้ช่องไฟบนมือถือ)
                # =========================================================
                st.markdown("<div style='height: 25px;'></div>", unsafe_allow_html=True)
                col_btn1, col_btn2, col_spacer, col_btn3 = st.columns([1.8, 1.5, 4.2, 2.5])
                with col_btn1: st.download_button("📗 Export เป็น Excel", data=excel_data_t1, file_name=f"Summary_Machine_Report_{start_date_t1}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="dl_t1")
                with col_btn2: 
                    if st.button("📊 สรุปรายงาน", use_container_width=True, key="view_t1"): show_summary_report_dialog(json.loads(json_html_t1))
                with col_btn3: components.html(f"""<body style="margin:0;padding:2px;overflow:hidden;"><button onclick="openPrintPreview1()" style="width:100%; height:42px; background-color:#0F172A; border:none; border-radius:24px; color:#F8FAFC; font-family:sans-serif; font-size:14.5px; font-weight:600; cursor:pointer; box-sizing:border-box; box-shadow: 0 4px 6px rgba(0,0,0,0.1); transition: all 0.2s ease;" onmouseover="this.style.backgroundColor='#1E293B'; this.style.transform='translateY(-1px)';" onmouseout="this.style.backgroundColor='#0F172A'; this.style.transform='translateY(0)';">🖨️ ปริ้นเอกสาร</button></body><script>function openPrintPreview1(){{var w = window.open('', '_blank', 'height=750,width=1100,scrollbars=yes'); w.document.write({json_html_t1}); w.document.close(); setTimeout(function(){{ w.print(); }}, 500);}}</script>""", height=50)
            
            else: st.info("ไม่พบข้อมูลรายงานตามช่วงเวลาที่เลือก")

    # =========================================================================
    # 🚚 TAB 2: รายงานสรุปการใช้เชื้อเพลิงรถ
    # =========================================================================
    elif key == "2":
        with current_tab_ctx:
            st.subheader("📊 รายงานสรุปการใช้เชื้อเพลิงรถยนต์และรถยก")
            col_f1, col_f2, col_f3 = st.columns([1.5, 1.5, 2])
            with col_f1: start_date_t2 = st.date_input("ตั้งแต่วันที่", format="DD/MM/YYYY", value=pd.to_datetime("today").replace(day=1), key="s_d2")
            with col_f2: end_date_t2 = st.date_input("ถึงวันที่", format="DD/MM/YYYY", value=pd.to_datetime("today"), key="e_d2")
            with col_f3:
                if current_role in ["admin", "reporter"]:
                    b_label_t2 = st.selectbox("เลือกสาขา:", options=report_options, key="b_s2")
                    b_id_t2 = "ทั้งหมด" if b_label_t2 == "ทั้งหมดทุกสาขา" else branch_dict[b_label_t2]
                    selected_branch_display = b_label_t2 if b_label_t2 != "ทั้งหมดทุกสาขา" else "ทุกสาขา"
                elif len(user_allowed_branches) > 1:
                    allowed_options = ["ทั้งหมดที่มีสิทธิ์"] + [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
                    b_label_t2 = st.selectbox("เลือกสาขา:", options=allowed_options, key="b_s2")
                    b_id_t2 = "รวมเฉพาะที่มีสิทธิ์" if b_label_t2 == "ทั้งหมดที่มีสิทธิ์" else branch_dict[b_label_t2]
                    selected_branch_display = b_label_t2 if b_label_t2 != "ทั้งหมดที่มีสิทธิ์" else "ทุกสาขาที่มีสิทธิ์"
                else:
                    st.info(f"📍 สังกัด: {user_branch_name}")
                    b_id_t2 = st.session_state.branch_id
                    selected_branch_display = user_branch_name

            raw_data_t2, excel_data_t2, json_html_t2, engine_lbl_list, engine_details_t2 = process_t2(b_id_t2, str(start_date_t2), str(end_date_t2), tuple(user_allowed_branches), selected_branch_display)

            if raw_data_t2:
                pk_col_t2 = list(raw_data_t2[0].keys())[0]
                paginated_t2 = render_custom_pagination(raw_data_t2, session_prefix="rep2")
                st.markdown("<br>", unsafe_allow_html=True)

                # =========================================================
                # 🖥️ DESKTOP VIEW: แสดงตารางแบบเดิม
                # =========================================================
                with st.container():
                    st.markdown('<div class="desktop-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                    # ปรับสัดส่วนคอลัมน์ใหม่ (ลบหมายเหตุออก เหลือ 8 คอลัมน์)
                    header_cols = st.columns([0.6, 1.2, 1.2, 1.2, 1.0, 1.0, 1.0, 1.2])
                    headers = ["No.", "วันที่", "สาขา/ประเภท", "ทะเบียนรถ", "ลิตร", "ชม.ทำงาน", "ลิตร/ชม.", "จัดการ"]
                    for col, header in zip(header_cols, headers): col.markdown(f"<span style='color:#64748B; font-weight:bold; font-size:14px;'>{header}</span>", unsafe_allow_html=True)
                    st.markdown("<hr style='margin: 0.2rem 0; border-color: #E2E8F0;'>", unsafe_allow_html=True)

                    with st.container(height=380, border=False):
                        start_num = (st.session_state.get("rep2_page", 1) - 1) * st.session_state.get("rep2_rows", 15)
                        for i, r in enumerate(paginated_t2):
                            seq_num = start_num + i + 1
                            rec_id = r[pk_col_t2]
                            lts, hrs = float(r.get('fuel_liters') or 0.0), float(r.get('working_hours') or 0.0)
                            cols = st.columns([0.6, 1.2, 1.2, 1.2, 1.0, 1.0, 1.0, 1.2])
                            cols[0].write(str(seq_num))
                            cols[1].write(pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-')
                            cols[2].write(f"{r.get('branch_name') or '-'} / {r.get('type_name') or '-'}")
                            cols[3].write(str(r.get('engine_code') or '-'))
                            cols[4].write(f"{lts:.2f}")
                            cols[5].write(f"{hrs:.2f}")
                            cols[6].write(f"{(round(lts/hrs, 2) if hrs > 0 else 0.00):.2f}")
                            # ลบ cols[7] หมายเหตุออก แล้วเลื่อนปุ่มจัดการมาแทน
                            
                            can_crud = current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)
                            if can_crud:
                                c_edit, c_del = cols[7].columns(2, gap="small")
                                if c_edit.button("✏️", key=f"e2_desk_{rec_id}", help="แก้ไขข้อมูล", use_container_width=True): update_record_dialog_t2(r, pk_col_t2, engine_lbl_list, engine_details_t2)
                                if c_del.button("🗑️", key=f"d2_desk_{rec_id}", help="ลบรายการ", use_container_width=True): delete_record_dialog_t2(r, pk_col_t2)
                            else: cols[7].write("-")
                            st.markdown("<hr style='margin:0; border-color:#F1F5F9;'>", unsafe_allow_html=True)

                # =========================================================
                # 📱 MOBILE VIEW: แสดงตารางแบบใหม่
                # =========================================================
                with st.container():
                    st.markdown('<div class="mobile-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                    if paginated_t2:
                        df_list = []
                        start_num = (st.session_state.get("rep2_page", 1) - 1) * st.session_state.get("rep2_rows", 15)
                        for i, r in enumerate(paginated_t2):
                            seq_num = start_num + i + 1
                            lts, hrs = float(r.get('fuel_liters') or 0.0), float(r.get('working_hours') or 0.0)
                            df_list.append({
                                "ลำดับ": seq_num,
                                "วันที่": pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-',
                                "สาขา/ประเภท": f"{r.get('branch_name') or '-'} / {r.get('type_name') or '-'}",
                                "ทะเบียนรถ": r.get('engine_code') or '-',
                                "ลิตร": f"{lts:.2f}",
                                "ชม.ทำงาน": f"{hrs:.2f}",
                                "ลิตร/ชม.": f"{(round(lts/hrs, 2) if hrs > 0 else 0.00):.2f}",
                                "หมายเหตุ": r.get('remark') or '-'
                            })
                        st.dataframe(pd.DataFrame(df_list), use_container_width=True, hide_index=True)
                        
                        st.markdown("---")
                        st.markdown("<h5 style='color:#333333; font-weight:bold;'>🛠️ จัดการข้อมูล (แก้ไข / ลบรายการ)</h5>", unsafe_allow_html=True)
                        editable_records = [r for r in paginated_t2 if current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)]
                        
                        if editable_records:
                            col_sel, col_btn_e, col_btn_d = st.columns([2.5, 1, 1])
                            opts = {f"ID: {r[pk_col_t2]} | วันที่: {pd.to_datetime(r.get('record_date')).strftime('%d/%m/%Y')} | {r.get('engine_code')}": r for r in editable_records}
                            sel_lbl = col_sel.selectbox("เลือกลำดับข้อมูล", options=list(opts.keys()), label_visibility="collapsed", key="sel_crud_t2_mob")
                            sel_row = opts[sel_lbl]
                            
                            if col_btn_e.button("✏️", use_container_width=True, key=f"e2_btn_mob"): update_record_dialog_t2(sel_row, pk_col_t2, engine_lbl_list, engine_details_t2)
                            if col_btn_d.button("🗑️", use_container_width=True, key=f"d2_btn_mob"): delete_record_dialog_t2(sel_row, pk_col_t2)
                        else:
                            st.info("ไม่มีสิทธิ์จัดการข้อมูล")

                # =========================================================
                # 🖨️ ปุ่มด้านล่าง
                # =========================================================
                st.markdown("<div style='height: 25px;'></div>", unsafe_allow_html=True)
                col_btn1, col_btn2, col_spacer, col_btn3 = st.columns([1.8, 1.5, 4.2, 2.5])
                with col_btn1: st.download_button("📗 Export เป็น Excel", data=excel_data_t2, file_name=f"Summary_Fuel_Report_{start_date_t2}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="dl_t2")
                with col_btn2: 
                    if st.button("📊 สรุปรายงาน", use_container_width=True, key="view_t2"): show_summary_report_dialog(json.loads(json_html_t2))
                with col_btn3: components.html(f"<body style='margin:0;padding:2px;overflow:hidden;'><button onclick='openPrintPreview2()' style='width:100%; height:42px; background-color:#0F172A; border:none; border-radius:24px; color:#F8FAFC; font-family:sans-serif; font-size:14.5px; font-weight:600; cursor:pointer; box-sizing:border-box; box-shadow: 0 4px 6px rgba(0,0,0,0.1); transition: all 0.2s ease;' onmouseover=\"this.style.backgroundColor='#1E293B'; this.style.transform='translateY(-1px)';\" onmouseout=\"this.style.backgroundColor='#0F172A'; this.style.transform='translateY(0)';\">🖨️ ปริ้นเอกสาร</button></body><script>function openPrintPreview2(){{var w = window.open('', '_blank', 'height=750,width=1100,scrollbars=yes'); w.document.write({json_html_t2}); w.document.close(); setTimeout(function(){{ w.print(); }}, 500);}}</script>", height=50)
            else: st.info("ไม่พบข้อมูลรายงานตามช่วงเวลาที่เลือก")

    # =========================================================================
    # 💨 TAB 3: รายงานสรุปแรงดันไอน้ำ บอยเลอร์
    # =========================================================================
    elif key == "3":
        with current_tab_ctx:
            st.subheader("📊 รายงานสรุปสถิติแรงดันไอน้ำปลายทางตก")
            col_f1, col_f2, col_f3 = st.columns([1.5, 1.5, 2])
            with col_f1: start_date_t3 = st.date_input("ตั้งแต่วันที่", format="DD/MM/YYYY", value=pd.to_datetime("today").replace(day=1), key="s_d3")
            with col_f2: end_date_t3 = st.date_input("ถึงวันที่", format="DD/MM/YYYY", value=pd.to_datetime("today"), key="e_d3")
            with col_f3:
                if current_role in ["admin", "reporter"]:
                    b_label_t3 = st.selectbox("เลือกสาขา:", options=report_options, key="b_s3")
                    b_id_t3 = "ทั้งหมด" if b_label_t3 == "ทั้งหมดทุกสาขา" else branch_dict[b_label_t3]
                    selected_branch_display = b_label_t3 if b_label_t3 != "ทั้งหมดทุกสาขา" else "ทุกสาขา"
                elif len(user_allowed_branches) > 1:
                    allowed_options = ["ทั้งหมดที่มีสิทธิ์"] + [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
                    b_label_t3 = st.selectbox("เลือกสาขา:", options=allowed_options, key="b_s3")
                    b_id_t3 = "รวมเฉพาะที่มีสิทธิ์" if b_label_t3 == "ทั้งหมดที่มีสิทธิ์" else branch_dict[b_label_t3]
                    selected_branch_display = b_label_t3 if b_label_t3 != "ทั้งหมดที่มีสิทธิ์" else "ทุกสาขาที่มีสิทธิ์"
                else:
                    st.info(f"📍 สังกัด: {user_branch_name}")
                    b_id_t3 = st.session_state.branch_id
                    selected_branch_display = user_branch_name

            raw_data_t3, excel_data_t3, json_html_t3 = process_t3(b_id_t3, str(start_date_t3), str(end_date_t3), tuple(user_allowed_branches), selected_branch_display)

            if raw_data_t3:
                pk_col_t3 = list(raw_data_t3[0].keys())[0]
                paginated_t3 = render_custom_pagination(raw_data_t3, session_prefix="rep3")
                st.markdown("<br>", unsafe_allow_html=True)

                # =========================================================
                # 🖥️ DESKTOP VIEW: แสดงตารางแบบเดิม
                # =========================================================
                with st.container():
                    st.markdown('<div class="desktop-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                    # ปรับสัดส่วนคอลัมน์ใหม่ (ลบหมายเหตุออก เหลือ 8 คอลัมน์)
                    header_cols = st.columns([0.6, 1.2, 1.0, 1.1, 1.1, 1.1, 1.1, 1.2])
                    headers = ["No.", "วันที่", "สาขา", "ทั้งหมด(ครั้ง)", "ตก(ครั้ง)", "ตก PM", "ตกนอก PM", "จัดการ"]
                    for col, header in zip(header_cols, headers): col.markdown(f"<span style='color:#64748B; font-weight:bold; font-size:14px;'>{header}</span>", unsafe_allow_html=True)
                    st.markdown("<hr style='margin: 0.2rem 0; border-color: #E2E8F0;'>", unsafe_allow_html=True)

                    with st.container(height=380, border=False):
                        start_num = (st.session_state.get("rep3_page", 1) - 1) * st.session_state.get("rep3_rows", 15)
                        for i, r in enumerate(paginated_t3):
                            seq_num = start_num + i + 1
                            rec_id = r[pk_col_t3]
                            cols = st.columns([0.6, 1.2, 1.0, 1.1, 1.1, 1.1, 1.1, 1.2])
                            cols[0].write(str(seq_num))
                            cols[1].write(pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-')
                            cols[2].markdown(f"<span class='branch-badge'>{r.get('branch_name') or '-'}</span>", unsafe_allow_html=True)
                            cols[3].write(str(r.get('total_count') or 0))
                            cols[4].write(str(r.get('total_drop') or 0))
                            cols[5].write(str(r.get('pm_drop') or 0))
                            cols[6].write(str(r.get('non_pm_drop') or 0))
                            # ลบ cols[7] หมายเหตุออก แล้วเลื่อนปุ่มจัดการมาแทน
                            
                            can_crud = current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)
                            if can_crud:
                                c_edit, c_del = cols[7].columns(2, gap="small")
                                if c_edit.button("✏️", key=f"e3_desk_{rec_id}", help="แก้ไขข้อมูล", use_container_width=True): update_record_dialog_t3(r, pk_col_t3)
                                if c_del.button("🗑️", key=f"d3_desk_{rec_id}", help="ลบรายการ", use_container_width=True): delete_record_dialog_t3(r, pk_col_t3)
                            else: cols[7].write("-")
                            st.markdown("<hr style='margin:0; border-color:#F1F5F9;'>", unsafe_allow_html=True)

                # =========================================================
                # 📱 MOBILE VIEW: แสดงตารางแบบใหม่
                # =========================================================
                with st.container():
                    st.markdown('<div class="mobile-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                    if paginated_t3:
                        df_list = []
                        start_num = (st.session_state.get("rep3_page", 1) - 1) * st.session_state.get("rep3_rows", 15)
                        for i, r in enumerate(paginated_t3):
                            seq_num = start_num + i + 1
                            df_list.append({
                                "ลำดับ": seq_num,
                                "วันที่": pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-',
                                "สาขา": r.get('branch_name') or '-',
                                "ทั้งหมด(ครั้ง)": r.get('total_count') or 0,
                                "ตก(ครั้ง)": r.get('total_drop') or 0,
                                "ตก PM": r.get('pm_drop') or 0,
                                "ตกนอก PM": r.get('non_pm_drop') or 0,
                                "หมายเหตุ": r.get('remark') or '-'
                            })
                        st.dataframe(pd.DataFrame(df_list), use_container_width=True, hide_index=True)

                        st.markdown("---")
                        st.markdown("<h5 style='color:#333333; font-weight:bold;'>🛠️ จัดการข้อมูล (แก้ไข / ลบรายการ)</h5>", unsafe_allow_html=True)
                        editable_records = [r for r in paginated_t3 if current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)]
                        
                        if editable_records:
                            col_sel, col_btn_e, col_btn_d = st.columns([2.5, 1, 1])
                            opts = {f"ID: {r[pk_col_t3]} | วันที่: {pd.to_datetime(r.get('record_date')).strftime('%d/%m/%Y')} | {r.get('branch_name')}": r for r in editable_records}
                            sel_lbl = col_sel.selectbox("เลือกลำดับข้อมูล", options=list(opts.keys()), label_visibility="collapsed", key="sel_crud_t3_mob")
                            sel_row = opts[sel_lbl]
                            
                            if col_btn_e.button("✏️", use_container_width=True, key=f"e3_btn_mob"): update_record_dialog_t3(sel_row, pk_col_t3)
                            if col_btn_d.button("🗑️", use_container_width=True, key=f"d3_btn_mob"): delete_record_dialog_t3(sel_row, pk_col_t3)
                        else:
                            st.info("ไม่มีสิทธิ์จัดการข้อมูล")

                # =========================================================
                # 🖨️ ปุ่มด้านล่าง
                # =========================================================
                st.markdown("<div style='height: 25px;'></div>", unsafe_allow_html=True)
                col_btn1, col_btn2, col_spacer, col_btn3 = st.columns([1.8, 1.5, 4.2, 2.5])
                with col_btn1: st.download_button("📗 Export เป็น Excel", data=excel_data_t3, file_name=f"Summary_Pressure_Report_{start_date_t3}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="dl_t3")
                with col_btn2: 
                    if st.button("📊 สรุปรายงาน", use_container_width=True, key="view_t3"): show_summary_report_dialog(json.loads(json_html_t3))
                with col_btn3: components.html(f"<body style='margin:0;padding:2px;overflow:hidden;'><button onclick='openPrintPreview3()' style='width:100%; height:42px; background-color:#0F172A; border:none; border-radius:24px; color:#F8FAFC; font-family:sans-serif; font-size:14.5px; font-weight:600; cursor:pointer; box-sizing:border-box; box-shadow: 0 4px 6px rgba(0,0,0,0.1); transition: all 0.2s ease;' onmouseover=\"this.style.backgroundColor='#1E293B'; this.style.transform='translateY(-1px)';\" onmouseout=\"this.style.backgroundColor='#0F172A'; this.style.transform='translateY(0)';\">🖨️ ปริ้นเอกสาร</button></body><script>function openPrintPreview3(){{var w = window.open('', '_blank', 'height=750,width=1100,scrollbars=yes'); w.document.write({json_html_t3}); w.document.close(); setTimeout(function(){{ w.print(); }}, 500);}}</script>", height=50)
            else: st.info("ไม่พบข้อมูลรายงานตามช่วงเวลาที่เลือก")

    # =========================================================================
    # 🔥 TAB 4: รายงานสรุปการใช้เชื้อเพลิงบอยเลอร์ (SYSTEM)
    # =========================================================================
    elif key == "4":
        with current_tab_ctx:
            st.subheader("📊 รายงานผลการคำนวณประสิทธิภาพเชื้อเพลิง บอยเลอร์")
            col_f1, col_f2, col_f3 = st.columns([1.5, 1.5, 2])
            with col_f1: start_date_t4 = st.date_input("ตั้งแต่วันที่", format="DD/MM/YYYY", value=pd.to_datetime("today").replace(day=1), key="s_d4")
            with col_f2: end_date_t4 = st.date_input("ถึงวันที่", format="DD/MM/YYYY", value=pd.to_datetime("today"), key="e_d4")
            with col_f3:
                if current_role in ["admin", "reporter"]:
                    b_label_t4 = st.selectbox("เลือกสาขา:", options=report_options, key="b_s4")
                    b_id_t4 = "ทั้งหมด" if b_label_t4 == "ทั้งหมดทุกสาขา" else branch_dict[b_label_t4]
                    selected_branch_display = b_label_t4 if b_label_t4 != "ทั้งหมดทุกสาขา" else "ทุกสาขา"
                elif len(user_allowed_branches) > 1:
                    allowed_options = ["ทั้งหมดที่มีสิทธิ์"] + [k for k, v in branch_dict.items() if str(v) in user_allowed_branches]
                    b_label_t4 = st.selectbox("เลือกสาขา:", options=allowed_options, key="b_s4")
                    b_id_t4 = "รวมเฉพาะที่มีสิทธิ์" if b_label_t4 == "ทั้งหมดที่มีสิทธิ์" else branch_dict[b_label_t4]
                    selected_branch_display = b_label_t4 if b_label_t4 != "ทั้งหมดที่มีสิทธิ์" else "ทุกสาขาที่มีสิทธิ์"
                else:
                    st.info(f"📍 สังกัด: {user_branch_name}")
                    b_id_t4 = st.session_state.branch_id
                    selected_branch_display = user_branch_name

            raw_data_t4, oven_raw_t4, excel_data_t4, json_html_t4 = process_t4(b_id_t4, str(start_date_t4), str(end_date_t4), tuple(user_allowed_branches), selected_branch_display)

            if raw_data_t4 or oven_raw_t4:
                pk_col_t4 = list(raw_data_t4[0].keys())[0] if raw_data_t4 else None
                pk_col_ov = list(oven_raw_t4[0].keys())[0] if oven_raw_t4 else None
                unique_boilers = sorted(list(set([str(r.get('boiler_name') or 'ไม่ระบุ') for r in raw_data_t4]))) if raw_data_t4 else []
                tab_names = ["📁 ข้อมูลยอดรวมสาขา (เตา/ไม้อบ/แรงดัน)"] + [f"🔥 ข้อมูล {b}" for b in unique_boilers]
                tabs_list = st.tabs(tab_names)

                # ==========================================
                # 🎯 แท็บย่อยที่ 0: ข้อมูลยอดรวมสาขา (เตา/ไม้อบ)
                # ==========================================
                with tabs_list[0]:
                    if oven_raw_t4:
                        paginated_ov = render_custom_pagination(oven_raw_t4, session_prefix="rep4_ov")
                        st.markdown("<br>", unsafe_allow_html=True)
                        
                        # 🖥️ DESKTOP VIEW
                        with st.container():
                            st.markdown('<div class="desktop-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                            header_cols = st.columns([0.5, 1.0, 1.0, 1.2, 1.2, 1.4, 1.1])
                            headers = ["No.", "วันที่", "สาขา", "จำนวนเตา(เตา)", "ไม้อบออก(ลบ.ฟ.)", "แรงดันเฉลี่ย", "จัดการ"]
                            for col, header in zip(header_cols, headers): col.markdown(f"<span style='color:#64748B; font-weight:bold; font-size:13px;'>{header}</span>", unsafe_allow_html=True)
                            st.markdown("<hr style='margin: 0.2rem 0; border-color: #E2E8F0;'>", unsafe_allow_html=True)

                            with st.container(height=380, border=False):
                                start_num = (st.session_state.get("rep4_ov_page", 1) - 1) * st.session_state.get("rep4_ov_rows", 15)
                                for i, r in enumerate(paginated_ov):
                                    seq_num = start_num + i + 1
                                    rec_id = r[pk_col_ov]
                                    cols = st.columns([0.5, 1.0, 1.0, 1.2, 1.2, 1.4, 1.1])
                                    cols[0].write(str(seq_num))
                                    cols[1].write(pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-')
                                    cols[2].markdown(f"<span class='branch-badge'>{r.get('branch_name') or '-'}</span>", unsafe_allow_html=True)
                                    cols[3].write(str(r.get('oven_qty') or 0))
                                    cols[4].write(f"{float(r.get('wood_out_cubft') or 0.0):.2f}")
                                    cols[5].write(f"{float(r.get('avg_terminal_pressure') or 0.0):.2f}")
                                    
                                    can_crud = current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)
                                    if can_crud:
                                        c_edit, c_del = cols[6].columns(2, gap="small")
                                        if c_edit.button("✏️", key=f"e_ov_desk_{rec_id}", help="แก้ไขข้อมูล", use_container_width=True): update_record_dialog_oven(r, pk_col_ov)
                                        if c_del.button("🗑️", key=f"d_ov_desk_{rec_id}", help="ลบรายการ", use_container_width=True): delete_record_dialog_oven(r, pk_col_ov)
                                    else: cols[6].write("-")
                                    st.markdown("<hr style='margin:0; border-color:#F1F5F9;'>", unsafe_allow_html=True)

                        # 📱 MOBILE VIEW
                        with st.container():
                            st.markdown('<div class="mobile-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                            if paginated_ov:
                                df_list = []
                                start_num = (st.session_state.get("rep4_ov_page", 1) - 1) * st.session_state.get("rep4_ov_rows", 15)
                                for i, r in enumerate(paginated_ov):
                                    seq_num = start_num + i + 1
                                    df_list.append({
                                        "ลำดับ": seq_num,
                                        "วันที่": pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-',
                                        "สาขา": r.get('branch_name') or '-',
                                        "จำนวนเตา(เตา)": r.get('oven_qty') or 0,
                                        "ไม้อบออก(ลบ.ฟ.)": f"{float(r.get('wood_out_cubft') or 0.0):.2f}",
                                        "แรงดันเฉลี่ย": f"{float(r.get('avg_terminal_pressure') or 0.0):.2f}"
                                    })
                                st.dataframe(pd.DataFrame(df_list), use_container_width=True, hide_index=True)

                                st.markdown("---")
                                st.markdown("<h5 style='color:#333333; font-weight:bold;'>🛠️ จัดการข้อมูลยอดรวมสาขา (แก้ไข / ลบ)</h5>", unsafe_allow_html=True)
                                editable_records = [r for r in paginated_ov if current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)]
                                
                                if editable_records:
                                    col_sel, col_btn_e, col_btn_d = st.columns([2.5, 1, 1])
                                    opts = {f"ID: {r[pk_col_ov]} | วันที่: {pd.to_datetime(r.get('record_date')).strftime('%d/%m/%Y')} | {r.get('branch_name')}": r for r in editable_records}
                                    sel_lbl = col_sel.selectbox("เลือกลำดับข้อมูล", options=list(opts.keys()), label_visibility="collapsed", key="sel_crud_ov_mob")
                                    sel_row = opts[sel_lbl]
                                    
                                    if col_btn_e.button("✏️", use_container_width=True, key=f"e_ov_btn_mob"): update_record_dialog_oven(sel_row, pk_col_ov)
                                    if col_btn_d.button("🗑️", use_container_width=True, key=f"d_ov_btn_mob"): delete_record_dialog_oven(sel_row, pk_col_ov)
                                else:
                                    st.info("ไม่มีสิทธิ์จัดการข้อมูล")
                    else:
                        st.info("ไม่พบข้อมูลยอดรวมสาขาในช่วงเวลาที่เลือก")

                # ==========================================
                # 🎯 แท็บย่อยอื่นๆ: ข้อมูลเชื้อเพลิงแยกตามบอยเลอร์
                # ==========================================
                def render_boiler_table(data_list, boiler_label):
                    if not data_list:
                        st.info(f"ไม่พบข้อมูลสำหรับ {boiler_label}")
                        return

                    session_prefix = f"rep4_b_{boiler_label}"
                    paginated_t4 = render_custom_pagination(data_list, session_prefix=session_prefix)
                    st.markdown("<br>", unsafe_allow_html=True)
                    
                    # 🖥️ DESKTOP VIEW
                    with st.container():
                        st.markdown('<div class="desktop-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                        header_cols = st.columns([0.5, 1.0, 0.8, 1.4, 1.3, 1.1, 1.4, 1.1])
                        headers = ["No.", "วันที่", "สาขา", "เชื้อเพลิงรวม(ตัน)", "ผลิตไอน้ำ(ตัน)", "ชม.ทำงาน", "ผลงาน(กก./ตัน)", "จัดการ"]
                        for col, header in zip(header_cols, headers): col.markdown(f"<span style='color:#64748B; font-weight:bold; font-size:13px;'>{header}</span>", unsafe_allow_html=True)
                        st.markdown("<hr style='margin: 0.2rem 0; border-color: #E2E8F0;'>", unsafe_allow_html=True)

                        with st.container(height=380, border=False):
                            start_num = (st.session_state.get(f"{session_prefix}_page", 1) - 1) * st.session_state.get(f"{session_prefix}_rows", 15)
                            for i, r in enumerate(paginated_t4):
                                seq_num = start_num + i + 1
                                rec_id = r[pk_col_t4]
                                s_w, w_w, ww_w = float(r.get('sawdust_weight') or 0.0), float(r.get('wood_weight') or 0.0), float(r.get('waste_wood_weight') or 0.0)
                                tot_w = s_w + w_w + ww_w
                                s_prod = float(r.get('steam_production') or 1.0)
                                cols = st.columns([0.5, 1.0, 0.8, 1.4, 1.3, 1.1, 1.4, 1.1])
                                cols[0].write(str(seq_num))
                                cols[1].write(pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-')
                                cols[2].markdown(f"<span class='branch-badge'>{r.get('branch_name') or '-'}</span>", unsafe_allow_html=True)
                                cols[3].write(f"{tot_w:.2f}"); cols[4].write(f"{s_prod:.2f}"); cols[5].write(f"{float(r.get('working_hours') or 0.0):.2f}"); cols[6].write(f"{((tot_w / s_prod) * 1000 if s_prod > 0 else 0.0):.2f}")
                                
                                can_crud = current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)
                                if can_crud:
                                    c_edit, c_del = cols[7].columns(2, gap="small")
                                    if c_edit.button("✏️", key=f"e4_{boiler_label}_desk_{rec_id}", help="แก้ไขข้อมูล", use_container_width=True): update_record_dialog_t4(r, pk_col_t4)
                                    if c_del.button("🗑️", key=f"d4_{boiler_label}_desk_{rec_id}", help="ลบรายการ", use_container_width=True): delete_record_dialog_t4(r, pk_col_t4)
                                else: cols[7].write("-")
                                st.markdown("<hr style='margin:0; border-color:#F1F5F9;'>", unsafe_allow_html=True)

                    # 📱 MOBILE VIEW
                    with st.container():
                        st.markdown('<div class="mobile-view-marker" style="display:none;"></div>', unsafe_allow_html=True)
                        if paginated_t4:
                            df_list = []
                            start_num = (st.session_state.get(f"{session_prefix}_page", 1) - 1) * st.session_state.get(f"{session_prefix}_rows", 15)
                            for i, r in enumerate(paginated_t4):
                                seq_num = start_num + i + 1
                                s_w, w_w, ww_w = float(r.get('sawdust_weight') or 0.0), float(r.get('wood_weight') or 0.0), float(r.get('waste_wood_weight') or 0.0)
                                tot_w = s_w + w_w + ww_w
                                s_prod = float(r.get('steam_production') or 1.0)
                                df_list.append({
                                    "ลำดับ": seq_num,
                                    "วันที่": pd.to_datetime(r.get('record_date')).strftime('%d-%m-%Y') if r.get('record_date') else '-',
                                    "สาขา": r.get('branch_name') or '-',
                                    "เชื้อเพลิงรวม(ตัน)": f"{tot_w:.2f}",
                                    "ผลิตไอน้ำ(ตัน)": f"{s_prod:.2f}",
                                    "ชม.ทำงาน": f"{float(r.get('working_hours') or 0.0):.2f}",
                                    "ผลงาน(กก./ตัน)": f"{((tot_w / s_prod) * 1000 if s_prod > 0 else 0.0):.2f}"
                                })
                            st.dataframe(pd.DataFrame(df_list), use_container_width=True, hide_index=True)

                            st.markdown("---")
                            st.markdown(f"<h5 style='color:#333333; font-weight:bold;'>🛠️ จัดการข้อมูล {boiler_label} (แก้ไข / ลบ)</h5>", unsafe_allow_html=True)
                            editable_records = [r for r in paginated_t4 if current_role == 'admin' or (current_role in ['user', 'manager'] and str(r.get('branch_id')) in user_allowed_branches)]
                            
                            if editable_records:
                                col_sel, col_btn_e, col_btn_d = st.columns([2.5, 1, 1])
                                opts = {f"ID: {r[pk_col_t4]} | วันที่: {pd.to_datetime(r.get('record_date')).strftime('%d/%m/%Y')} | {r.get('branch_name')}": r for r in editable_records}
                                sel_lbl = col_sel.selectbox("เลือกลำดับข้อมูล", options=list(opts.keys()), label_visibility="collapsed", key=f"sel_crud_{boiler_label}_mob")
                                sel_row = opts[sel_lbl]
                                if col_btn_e.button("✏️", use_container_width=True, key=f"e_{boiler_label}_btn_mob"): update_record_dialog_t4(sel_row, pk_col_t4)
                                if col_btn_d.button("🗑️", use_container_width=True, key=f"d_{boiler_label}_btn_mob"): delete_record_dialog_t4(sel_row, pk_col_t4)
                            else:
                                st.info("ไม่มีสิทธิ์จัดการข้อมูล")

                # ==========================================
                for idx, boiler_label in enumerate(unique_boilers):
                    with tabs_list[idx + 1]:
                        data_boiler = [r for r in raw_data_t4 if str(r.get('boiler_name') or 'ไม่ระบุ') == boiler_label]
                        render_boiler_table(data_boiler, boiler_label)

                # ==========================================
                # 🖨️ ปุ่มด้านล่าง (ใส่กล่องคั่นระยะ 25px แก้ปุ่มเบียด)
                # ==========================================
                st.markdown("<div style='height: 25px;'></div>", unsafe_allow_html=True)
                col_btn1, col_btn2, col_spacer, col_btn3 = st.columns([1.8, 1.5, 4.2, 2.5])
                with col_btn1: st.download_button("📗 Export เป็น Excel", data=excel_data_t4, file_name=f"Report_Boiler_Fuel_{start_date_t4}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="dl_t4")
                with col_btn2: 
                    if st.button("📊 สรุปรายงาน", use_container_width=True, key="view_t4"): 
                        parsed_data = json.loads(json_html_t4)
                        show_summary_report_dialog_t4(parsed_data["boilers"])
                with col_btn3: 
                    parsed_print_data = json.loads(json_html_t4)
                    boilers_print_dict = parsed_print_data.get("boilers", {})
                    master_html_str = ""
                    for k, v in boilers_print_dict.items():
                        if "รวมสาขา" in k or "รวม 1+2" in k:
                            master_html_str = v
                            break
                    boiler_html_snippets = {}
                    for b_lbl in unique_boilers:
                        found_html = ""
                        for k, v in boilers_print_dict.items():
                            if b_lbl in k and "รวมสาขา" not in k and "รวม 1+2" not in k:
                                found_html = v
                                break
                        boiler_html_snippets[b_lbl] = found_html

                    components.html(f"""
                    <body style="margin:0;padding:2px;overflow:hidden;">
                        <button onclick="openPrintPreview4()" style="width:100%; height:42px; background-color:#0F172A; border:none; border-radius:24px; color:#F8FAFC; font-family:sans-serif; font-size:14.5px; font-weight:600; cursor:pointer; box-sizing:border-box; box-shadow: 0 4px 6px rgba(0,0,0,0.1); transition: all 0.2s ease;" onmouseover="this.style.backgroundColor='#1E293B'; this.style.transform='translateY(-1px)';" onmouseout="this.style.backgroundColor='#0F172A'; this.style.transform='translateY(0)';">🖨️ ปริ้นเอกสาร</button>
                    </body>
                    <script>
                    function openPrintPreview4() {{
                        try {{
                            var activeTabs = window.parent.document.querySelectorAll('[role="tab"][aria-selected="true"], [data-baseweb="tab"][aria-selected="true"]');
                            var selectedTabName = "";
                            for(var i=0; i<activeTabs.length; i++) {{
                                var text = activeTabs[i].innerText || activeTabs[i].textContent;
                                if(text.includes("ยอดรวมสาขา") || text.includes("ข้อมูล บอยเลอร์") || text.includes("ข้อมูลบอยเลอร์")) {{
                                    selectedTabName = text.trim();
                                }}
                            }}
                            var printHtml = "";
                            var masterHtml = {json.dumps(master_html_str)};
                            var boilerSnippets = {json.dumps(boiler_html_snippets)};

                            if(selectedTabName.includes("ยอดรวมสาขา") || selectedTabName === "") {{ printHtml = masterHtml; }} 
                            else {{
                                for (var bKey in boilerSnippets) {{
                                    if(selectedTabName.includes(bKey)) {{ printHtml = boilerSnippets[bKey]; break; }}
                                }}
                                if(printHtml === "") {{ printHtml = masterHtml; }}
                            }}
                            var w = window.open('', '_blank', 'height=750,width=1100,scrollbars=yes');
                            w.document.write(printHtml); w.document.close(); setTimeout(function(){{ w.print(); }}, 500);
                        }} catch (e) {{
                            var w = window.open('', '_blank', 'height=750,width=1100,scrollbars=yes');
                            w.document.write({json.dumps(master_html_str)}); w.document.close(); setTimeout(function(){{ w.print(); }}, 500);
                        }}
                    }}
                    </script>
                    """, height=50)
            else: st.info("ไม่พบข้อมูลรายงานตามช่วงเวลาที่เลือก")